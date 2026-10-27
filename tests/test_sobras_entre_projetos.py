"""Sobra de um projeto do lote usada no seguinte: cores e observações automáticas na capa."""
import openpyxl

from retalhos.models import Chapa
from retalhos.services.capa_lote import CORES_SOBRAS, capa_de_analise, gerar_capa
from retalhos.services.consumo import analisar, exportar_excel, ler_plano
from tests.test_consumo import P, _peca, _plano, _program

CHAPAS = {10: Chapa(10, "15", "BRANCO TX"), 20: Chapa(20, "18", "CUMARU")}


def _sobra(codigo: int, largura: float, altura: float) -> str:
    """Etiqueta de sobra (template 2) como o Corte Certo grava: código de barras de 21 caracteres."""
    barras = f"{codigo:07d}0{largura:06.1f}{altura:07.1f}"
    assert len(barras) == 21
    return (f'<{P}material.program.cut type="X"><{P}material.program.cut.data template="2">'
            f'<{P}material.program.cut.data.field name="160" value="{barras}"/>'
            f'</{P}material.program.cut.data></{P}material.program.cut>')


def _lote_em_cadeia(tmp_path):
    """Projeto 100 gera sobras usadas pelo 200; o 200 gera sobra usada pelo 300."""
    a = _plano(tmp_path, "100_a.cutplanning", {
        10: [_program(False, 2740, 1840, 1, (_peca("100"), _sobra(10, 700, 600)))],
        20: [_program(False, 2740, 1840, 1, (_peca("100"), _sobra(20, 900, 500)))],
    })
    b = _plano(tmp_path, "200_b.cutplanning", {
        10: [_program(True, 600, 700, 1, (_peca("200"),)),  # sobra do 100 (medida invertida)
             _program(False, 2740, 1840, 1, (_peca("200"), _sobra(10, 800, 400)))],
    })
    c = _plano(tmp_path, "300_c.cutplanning", {
        10: [_program(True, 800, 400, 1, (_peca("300"),)),  # sobra do 200
             _program(True, 800, 400, 1, (_peca("300"),))],  # 2º igual: não há outra sobra -> estoque
        20: [_program(True, 900, 500, 1, (_peca("300"),))],  # sobra do 100
    })
    return [ler_plano(x) for x in (a, b, c)]


def test_sobras_lidas_do_plano(tmp_path):
    planos = _lote_em_cadeia(tmp_path)
    assert [(s.codigo, s.largura, s.altura) for s in planos[0].sobras] == [(10, 700, 600), (20, 900, 500)]
    assert len(planos[1].sobras) == 1 and planos[2].sobras == []


def test_passagens_e_origem_dos_retalhos(tmp_path):
    analise = analisar(_lote_em_cadeia(tmp_path), [], CHAPAS)
    assert analise.passagens == [("100", "200"), ("100", "300"), ("200", "300")]
    origens = sorted((r.projeto, r.codigo, r.medida, r.origem) for r in analise.retalhos)
    assert origens == [("200", 10, "600 x 700", "100"),
                       ("300", 10, "800 x 400", ""),      # cada sobra vira um retalho só
                       ("300", 10, "800 x 400", "200"),
                       ("300", 20, "900 x 500", "100")]


def _plano_a_e_b(tmp_path, usos_no_b: int):
    a = _plano(tmp_path, "100_a.cutplanning", {
        10: [_program(False, 2740, 1840, 1, (_peca("100"), _sobra(10, 700, 600)))],
    })
    b = _plano(tmp_path, "200_b.cutplanning", {10: [_program(True, 700, 600, usos_no_b, (_peca("200"),))]})
    return [ler_plano(a), ler_plano(b)]


def test_retalho_com_setor_nunca_e_sobra(tmp_path):
    from retalhos.models import Retalho
    guardado = Retalho(1, "5", 10, "15", "BRANCO TX", 700, 600, None, "", "", "", True)  # mesma medida, setor 5
    analise = analisar(_plano_a_e_b(tmp_path, usos_no_b=1), [guardado], CHAPAS)
    r = analise.retalhos[0]
    assert (r.retalho.id, r.setor, r.origem) == (1, "5", "")  # estoque, não sobra
    assert analise.passagens == []


def test_retalho_com_setor_nao_consome_a_sobra(tmp_path):
    from retalhos.models import Retalho
    guardado = Retalho(1, "5", 10, "15", "BRANCO TX", 700, 600, None, "", "", "", True)
    # o projeto 200 usa 2 retalhos de 700 x 600: um é o guardado no setor 5, o outro é a sobra do 100
    analise = analisar(_plano_a_e_b(tmp_path, usos_no_b=2), [guardado], CHAPAS)
    resumo = sorted(((r.retalho.id if r.retalho else None, r.origem) for r in analise.retalhos), key=str)
    assert resumo == sorted([(1, ""), (None, "100")], key=str)
    assert analise.passagens == [("100", "200")]


def test_sobra_lancada_sem_setor_continua_sendo_sobra(tmp_path):
    from retalhos.models import Retalho
    # a sobra do 100 já foi lançada pelo Leitor XML: sem setor, reservada para o 200
    lancada = Retalho(7, "", 10, "15", "BRANCO TX", 700, 600, None, "200", "", "", True)
    guardado = Retalho(1, "5", 10, "15", "BRANCO TX", 700, 600, None, "", "", "", True)
    analise = analisar(_plano_a_e_b(tmp_path, usos_no_b=1), [guardado, lancada], CHAPAS)
    r = analise.retalhos[0]
    assert (r.retalho.id, r.origem) == (7, "100")  # a reservada p/ o projeto tem prioridade e não tem setor


def test_sobra_do_proprio_plano_nao_conta(tmp_path):
    plano = ler_plano(_plano(tmp_path, "1_x.cutplanning", {
        10: [_program(True, 700, 600, 1, (_peca("1"),)), _program(False, 2740, 1840, 1, (_sobra(10, 700, 600),))],
    }))
    assert analisar([plano], [], CHAPAS).passagens == []


def test_capa_pinta_linhas_e_preenche_observacoes(tmp_path):
    analise = analisar(_lote_em_cadeia(tmp_path), [], CHAPAS)
    capa = capa_de_analise(analise, observacoes="Conferir furação.")
    assert capa.observacoes.splitlines() == [
        "Retalho AMARELO: sobra do projeto 100 p/ 200",
        "Retalho VERDE: sobra do projeto 100 p/ 300",
        "Retalho AZUL: sobra do projeto 200 p/ 300",
        "Conferir furação.",
    ]
    ws = openpyxl.load_workbook(gerar_capa(capa, tmp_path / "capa.xlsx"))["Frente"]
    cab = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Setor")
    fundos = sorted((ws.cell(r, 10).value, ws.cell(r, 8).value, ws.cell(r, 3).fill.fgColor.rgb[-6:])
                    for r in range(cab + 1, cab + 5))
    amarelo, verde, azul = (c for _, c in CORES_SOBRAS[:3])
    vermelho = "FDE2E2"  # o 2º retalho 800 x 400 não é sobra e não está no estoque
    assert fundos == sorted([("200", "600 x 700", amarelo), ("300", "800 x 400", azul),
                             ("300", "800 x 400", vermelho), ("300", "900 x 500", verde)])
    # a explicação das cores aparece uma vez só: nas Observações (sem legenda repetida abaixo da tabela)
    textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    explicacoes = [t for t in textos if t.lower().startswith("retalho amarelo")]
    assert len(explicacoes) == 1  # só a caixa de Observações
    assert explicacoes[0].splitlines()[0] == "Retalho AMARELO: sobra do projeto 100 p/ 200"


def test_excel_tem_coluna_de_origem(tmp_path):
    analise = analisar(_lote_em_cadeia(tmp_path), [], CHAPAS)
    ws = openpyxl.load_workbook(exportar_excel(analise, tmp_path / "c.xlsx"))["Retalhos usados"]
    assert ws.cell(1, 8).value == "Sobra do projeto"
    assert {ws.cell(r, 8).value for r in range(2, 6)} == {"100", "200", None}
