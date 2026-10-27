from datetime import datetime

import openpyxl

from retalhos.models import Chapa, Retalho
from retalhos.services import EstoqueService
from retalhos.services.capa_lote import (
    ESCALA_MINIMA_LEGIVEL, CapaLote, ChapaLote, ProjetoLote, RetalhoLote, avaliar_impressao, capa_de_analise,
    gerar_capa, ler_conf_folha_lote,
)
from retalhos.services.consumo import analisar, ler_plano
from tests.test_consumo import _peca, _plano, _program


def _conf_folha(path, linhas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lote 2566"
    ws.append(["C.C.", "Pedido", "Ambiente", "Cliente", "Entrega", "Lote"])
    for linha in linhas:
        ws.append(linha)
    wb.save(path)
    return path


def test_ler_conf_folha_lote(tmp_path):
    arq = _conf_folha(tmp_path / "conf.xlsx", [
        ["13955", "5479", "SOLTEIRO", "GABRIELA", "18/11/2026", "2566"],
        ["13956", "5480", "BANHO", "GABRIELA", "20/11/2026", "2566"],
        [None, None, None, None, None, None],
    ])
    lote, projetos = ler_conf_folha_lote(arq)
    assert lote == "2566"
    assert [(p.projeto, p.ambiente, p.entrega) for p in projetos] == [
        ("13955", "SOLTEIRO", "18/11/2026"), ("13956", "BANHO", "20/11/2026"),
    ]


def test_capa_de_analise_completa_projetos(tmp_path):
    plano = ler_plano(_plano(tmp_path, "13955_x.cutplanning", {
        10: [_program(False, 2740, 1840, 2), _program(True, 1500, 900), _program(True, 700, 600)],
    }))
    estoque = [Retalho(5, "CORTE", 10, "15", "BRANCO TX", 1500, 900, None, "", "", "", True)]
    analise = analisar([plano], estoque, {10: Chapa(10, "15", "BRANCO TX")})
    capa = capa_de_analise(analise, [ProjetoLote("13955", "5479", "SOLTEIRO", "GABRIELA", "18/11/2026")], "2566")
    assert capa.lote == "2566" and capa.projetos[0].cliente == "GABRIELA"
    assert [(c.espessura, c.quantidade, c.medida) for c in capa.chapas] == [("15", 2, "2740 x 1840")]
    assert [(r.setor, r.encontrado) for r in capa.retalhos] == [("CORTE", True), ("", False)]


def test_gerar_capa(tmp_path):
    capa = CapaLote(
        "2561", datetime(2026, 10, 17, 14, 6),
        projetos=[ProjetoLote("13929", "5441", "COPA SUPERIOR", "ALAN", "09/11/2026"),
                  ProjetoLote("13930", "5442", "COZINHA", "ALAN", "05/11/2026")],
        chapas=[ChapaLote("BRANCO TX", "15", 5, "2740 x 1840"), ChapaLote("BRANCO TX", "6", 2, "2740 x 1840"),
                ChapaLote("CUMARU", "18", 1, "2740 x 1840"), ChapaLote("CARVALHO", "25", 1, "2750 x 1850"),
                ChapaLote("BRANCO TX", "10", 1, "2740 x 1840")],
        retalhos=[RetalhoLote("ESTOQUE", "15", "BEIGE MATT", "2740 x 1840", "13929")],
        observacoes="Teste de observação",
    )
    assert capa.total_chapas == 10 and capa.proxima_entrega == "05/11/2026"
    ws = openpyxl.load_workbook(gerar_capa(capa, tmp_path / "capa" / "LOTE_2561.xlsx")).active
    textos = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "LOTE 2561" in textos and "Teste de observação" in textos
    assert {"Material", "Esp.", "Medida", "Qtd", "TOTAL", "6 mm", "10 mm", "15 mm", "18 mm", "25 mm"} <= set(textos)
    # lista: material ordenado e espessura em ordem numérica (6, 10, 15 - não "10" antes de "6")
    linhas = [(ws.cell(r, 1).value, ws.cell(r, 7).value) for r in range(1, ws.max_row + 1)
              if ws.cell(r, 7).value in ("6 mm", "10 mm", "15 mm", "18 mm", "25 mm")]
    assert linhas == [("BRANCO TX", "6 mm"), ("BRANCO TX", "10 mm"), ("BRANCO TX", "15 mm"),
                      ("CARVALHO", "25 mm"), ("CUMARU", "18 mm")]
    assert ws.page_setup.orientation == "portrait" and ws.sheet_properties.pageSetUpPr.fitToPage
    # sem macros, sem fórmulas e sem vínculos externos
    assert not any(isinstance(v, str) and v.startswith("=") for v in textos)


def test_capa_lote_com_varios_projetos(tmp_path):
    pecas = [_peca(p, "GABRIELA", f"{5479 + i} - AMB {i}") for i, p in enumerate(["13955", "13956", "13957"])]
    pecas.append(_peca("13990", "JOAO", "6000 - ESCRITORIO"))
    plano = ler_plano(_plano(tmp_path, "13955_GABRIELA.cutplanning", {
        10: [_program(False, 2740, 1840, 3, tuple(pecas))],
    }))
    analise = analisar([plano], [], {10: Chapa(10, "15", "BRANCO TX")})
    conf = [ProjetoLote("13955", "5479", "SOLTEIRO", "GABRIELA", "18/11/2026"),
            ProjetoLote("13956", "5480", "BANHO SUITE", "GABRIELA", "18/11/2026"),
            ProjetoLote("13957", "5481", "BANHO SOCIAL", "GABRIELA", "20/11/2026")]
    capa = capa_de_analise(analise, conf, "2566")

    assert [p.projeto for p in capa.projetos] == ["13955", "13956", "13957", "13990"]
    fora = capa.projetos[3]
    assert (fora.no_conf_folha, fora.pedido, fora.ambiente, fora.cliente, fora.entrega) == (
        False, "6000", "ESCRITORIO", "JOAO", "")
    assert capa.projetos[0].ambiente == "SOLTEIRO"  # ConFolhaLote tem prioridade sobre a etiqueta

    wb = openpyxl.load_workbook(gerar_capa(capa, tmp_path / "LOTE_2566.xlsx"))
    textos = [c.value for row in wb["Verso"].iter_rows() for c in row if isinstance(c.value, str)]
    cartoes = [t for t in textos if t.startswith("Projeto 1")]
    assert len(cartoes) == 4 and "Projeto 13990  ·  ESCRITORIO  ·  JOAO" in cartoes


def _capa_grande(n_projetos=9, n_chapas=25, n_retalhos=60) -> CapaLote:
    return CapaLote(
        "1", datetime(2026, 1, 1),
        projetos=[ProjetoLote(str(14000 + i), cliente=f"C{i}") for i in range(n_projetos)],
        chapas=[ChapaLote(f"MATERIAL {i:02}", "15", 1, "2740 x 1840") for i in range(n_chapas)],
        retalhos=[RetalhoLote("CORTE", "15", f"MAT {i:02}", f"{500 + i} x 400", "14000") for i in range(n_retalhos)],
    )


def test_frente_e_verso_sempre_uma_pagina_cada(tmp_path):
    for capa in (CapaLote("1", datetime(2026, 1, 1), [ProjetoLote("1")]), _capa_grande()):
        wb = openpyxl.load_workbook(gerar_capa(capa, tmp_path / "c.xlsx"))
        assert wb.sheetnames == ["Frente", "Verso"] and wb.active.title == "Frente"
        for ws in wb.worksheets:
            setup = ws.page_setup
            assert (setup.fitToWidth, setup.fitToHeight, ws.sheet_properties.pageSetUpPr.fitToPage) == (1, 1, True)
            assert not ws.row_breaks.brk  # nenhuma quebra manual: 1 página por face
        verso = [c.value for row in wb["Verso"].iter_rows() for c in row if isinstance(c.value, str)]
        assert any(t.startswith("CONTROLE DE MONTAGEM") for t in verso)
        frente = [c.value for row in wb["Frente"].iter_rows() for c in row if isinstance(c.value, str)]
        assert not any(t.startswith("CONTROLE DE MONTAGEM") for t in frente)


def test_modo_compacto_e_aviso_de_escala():
    pequena = avaliar_impressao(CapaLote("1", datetime(2026, 1, 1), [ProjetoLote(str(i)) for i in range(4)]))
    assert not pequena.frente_compacta and not pequena.verso_compacto and not pequena.avisos

    media = avaliar_impressao(_capa_grande(n_projetos=6, n_chapas=12, n_retalhos=25))
    assert media.frente_compacta and media.verso_compacto  # coube encolhendo linhas, não a letra
    assert media.escala_frente >= ESCALA_MINIMA_LEGIVEL and not media.avisos

    grande = avaliar_impressao(_capa_grande(n_projetos=12, n_chapas=40, n_retalhos=90))
    assert grande.escala_frente < ESCALA_MINIMA_LEGIVEL
    assert any("frente" in a for a in grande.avisos)


def test_producao_com_linha_reduzida(tmp_path):
    wb = openpyxl.load_workbook(gerar_capa(CapaLote("1", datetime(2026, 1, 1), [ProjetoLote("777")]),
                                           tmp_path / "c.xlsx"))
    ws = wb["Frente"]
    linha = next(c.row for row in ws.iter_rows() for c in row if c.value == "777")
    assert ws.row_dimensions[linha].height == 18


def test_retalhos_um_por_linha_sem_qtd_e_material_largo(tmp_path):
    plano = ler_plano(_plano(tmp_path, "5_x.cutplanning", {
        10: [_program(True, 700, 600, 3), _program(True, 900, 600)],
    }))
    estoque = [Retalho(1, "7", 10, "15", "BRANCO TX", 700, 600, None, "", "", "", True)]
    capa = capa_de_analise(analisar([plano], estoque, {10: Chapa(10, "15", "BRANCO TX")}))
    assert len(capa.retalhos) == 4 and capa.total_retalhos == 4  # cada retalho é único: sem agrupar

    ws = openpyxl.load_workbook(gerar_capa(capa, tmp_path / "c.xlsx"))["Frente"]
    cab = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Setor")
    titulos = [ws.cell(cab, c).value for c in range(1, 11) if ws.cell(cab, c).value]
    assert titulos == ["Setor", "Esp.", "Material", "Medida", "Projeto"]  # sem "Qtd"
    faixa_material = next(m for m in ws.merged_cells.ranges if m.min_row == cab and m.min_col == 3)
    assert (faixa_material.min_col, faixa_material.max_col) == (3, 7)  # material com 5 colunas da grade
    assert ws.column_dimensions["A"].width == 9  # setor estreito
    linhas = [[ws.cell(r, c).value for c in (1, 3)] for r in range(cab + 1, cab + 5)]
    assert ["7", "BRANCO TX"] in linhas and linhas.count(["—", "BRANCO TX"]) == 3
    assert any(isinstance(c.value, str) and c.value.startswith("Linhas em vermelho") for row in ws.iter_rows() for c in row)



def test_service_capa_do_lote(settings, tmp_path):
    settings.conf_folha_lote = str(_conf_folha(tmp_path / "conf.xlsx",
                                               [["300", "1", "SALA", "MARIA", "01/12/2026", "2600"]]))
    svc = EstoqueService(settings)
    arquivo = _plano(tmp_path, "300_x.cutplanning", {10: [_program(False, 2740, 1840)]})
    capa = svc.capa_do_lote(svc.analisar_planos([arquivo]))
    assert capa.lote == "2600" and capa.projetos[0].cliente == "MARIA"

    # projeto fora do ConFolhaLote: lote não é sugerido; planilha ausente não é erro
    outro = _plano(tmp_path, "999_y.cutplanning", {10: [_program(False, 2740, 1840)]})
    assert svc.capa_do_lote(svc.analisar_planos([outro])).lote == ""
    settings.conf_folha_lote = str(tmp_path / "nao_existe.xlsx")
    assert svc.capa_do_lote(svc.analisar_planos([outro])).projetos[0].projeto == "999"
