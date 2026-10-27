"""Consulta de retalhos removidos (baixa/exclusão), por material, e restauração pelo administrador."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from retalhos.models import Usuario  # noqa: E402
from retalhos.services import EstoqueService, RegraDeNegocioError  # noqa: E402


def _por_id(svc, rid):
    return next(r for r in svc.listar_retalhos() if r.id == rid)


def _cenario(settings):
    settings.usuario = Usuario(1, "ADM", "ANA", admin=True)
    svc = EstoqueService(settings)
    baixado = svc.lancar_retalho(10, 1000, 500)
    excluido = svc.lancar_retalho(10, 800, 400)
    outro_material = svc.lancar_retalho(20, 700, 400)
    ativo = svc.lancar_retalho(10, 600, 400)
    svc.reservar("1206", [_por_id(svc, baixado)], gerar_planilha=False)
    svc.dar_baixa("1206")
    svc.excluir_retalho(_por_id(svc, excluido))
    svc.excluir_retalho(_por_id(svc, outro_material))
    return svc, baixado, excluido, outro_material, ativo


def test_lista_por_material(settings):
    svc, baixado, excluido, outro_material, ativo = _cenario(settings)
    assert svc.contar_removidos_por_material() == {10: 2, 20: 1}

    removidos = svc.listar_removidos(10)
    assert {x.retalho.id for x in removidos} == {baixado, excluido}  # só do material pedido, sem os ativos
    por_id = {x.retalho.id: x for x in removidos}
    assert por_id[baixado].motivo == "Baixa" and por_id[baixado].retalho.reserva == "1206"
    assert por_id[excluido].motivo == "Exclusão"
    assert all(x.responsavel == "ANA" and x.data is not None for x in removidos)
    assert [x.retalho.id for x in svc.listar_removidos(20)] == [outro_material]


def test_restaurar_volta_disponivel_e_registra(settings):
    svc, baixado, excluido, _outro, ativo = _cenario(settings)
    removido = next(x for x in svc.listar_removidos(10) if x.retalho.id == baixado)
    svc.restaurar_retalho(removido)

    r = _por_id(svc, baixado)
    assert not r.reservado and r.responsavel == "ANA"  # disponível; lançamento original mantido
    assert baixado not in {x.retalho.id for x in svc.listar_removidos(10)}
    ret = svc.arquivos.arquivo_retalho(10).read_text()
    assert str(baixado) in [linha.split(",")[0] for linha in ret.splitlines()]
    ultimo = svc.historico(baixado)[-1]
    assert ultimo.acao == "RESTAURAÇÃO" and "baixado no projeto 1206" in ultimo.detalhe

    with pytest.raises(RegraDeNegocioError, match="já voltou"):
        svc.restaurar_retalho(removido)  # outra tela já restaurou

    excl = next(x for x in svc.listar_removidos(10) if x.retalho.id == excluido)
    svc.restaurar_retalho(excl)
    assert "estava excluído" in svc.historico(excluido)[-1].detalhe


def test_restaurar_somente_admin(settings):
    svc, *_ = _cenario(settings)
    removido = svc.listar_removidos(10)[0]
    settings.usuario = Usuario(2, "PROD", "PRODUCAO")
    with pytest.raises(RegraDeNegocioError, match="administradores"):
        svc.restaurar_retalho(removido)


def test_exportar(settings, tmp_path):
    svc, *_ = _cenario(settings)
    arquivo = svc.exportar_removidos(svc.listar_removidos(10), tmp_path / "rem.xlsx")
    ws = load_workbook(arquivo).active
    assert ws.max_row == 3 and {ws.cell(r, 8).value for r in (2, 3)} == {"Baixa", "Exclusão"}


def test_tela(settings):
    from retalhos.app import create_app
    from retalhos.ui.dialogs import RemovidosDialog
    create_app()
    svc, baixado, *_ = _cenario(settings)
    dlg = RemovidosDialog(svc)
    assert dlg.table.source.rows() == []  # nada é listado antes de escolher o material
    assert dlg.cb_material.count() == 2 and dlg.cb_material.currentIndex() == -1
    dlg.cb_material.setCurrentIndex(dlg.cb_material.findData(10))
    dlg._carregar()
    assert len(dlg.table.visible_rows()) == 2
    dlg.cb_motivo.setCurrentText("Baixas")
    assert [x.retalho.id for x in dlg.table.visible_rows()] == [baixado]
    assert dlg.bt_restaurar.isVisibleTo(dlg)

    settings.usuario = Usuario(2, "PROD", "PRODUCAO")
    outro = RemovidosDialog(svc)
    assert not outro.bt_restaurar.isVisibleTo(outro)
