from pathlib import Path

import openpyxl
import pytest

from retalhos.models import Chapa, Retalho, fmt_number
from retalhos.services.consumo import (
    SITUACAO_DISPONIVEL, SITUACAO_NAO_ENCONTRADO, SITUACAO_OUTRO, SITUACAO_PROJETO, analisar,
    exportar_excel, ler_plano, projeto_do_arquivo, resumo_chapas, resumo_retalhos,
)
from retalhos.services import EstoqueService

DADOS = Path(__file__).resolve().parent / "dados"  # planos de corte reais para teste
P = "com.geeksystem.cutplanning.cutplan."


def _peca(projeto: str, cliente: str = "CLIENTE", ambiente: str = "5000 - COZINHA") -> str:
    """Etiqueta de peça (template 1) como o Corte Certo grava: projeto no nome da imagem (campo 172)."""
    campos = {"172": f"projects/{cliente}-images/\\PLA{projeto}_00001.bmp", "184": cliente, "187": ambiente}
    fields = "".join(f'<{P}material.program.cut.data.field name="{k}" value="{v}"/>' for k, v in campos.items())
    return (f'<{P}material.program.cut type="X"><{P}material.program.cut.data template="1">{fields}'
            f'</{P}material.program.cut.data></{P}material.program.cut>')


def _program(iswaste: bool, lenght: float, width: float, qtd: int = 1, pecas: tuple[str, ...] = ()) -> str:
    return (f'<{P}material.program iswaste="{str(iswaste).lower()}" lenght="{lenght}" width="{width}" '
            f'quantity="{qtd}">{"".join(pecas)}</{P}material.program>')


def _plano(tmp_path: Path, nome: str, materiais: dict[int, list[str]]) -> Path:
    corpo = "".join(
        f'<{P}material id="{cod}" description="MDF {cod}" thickness="15.0">{"".join(progs)}</{P}material>'
        for cod, progs in materiais.items()
    )
    arquivo = tmp_path / nome
    arquivo.write_text(f'<?xml version="1.0" encoding="UTF-8"?><{P[:-1]} code="X">{corpo}</{P[:-1]}>',
                       encoding="utf-8")
    return arquivo


def _ret(rid, cod, larg, alt, reserva=""):
    return Retalho(rid, "CORTE", cod, "15", f"MAT {cod}", larg, alt, None, reserva, "", "", True)


def test_projeto_do_arquivo():
    assert projeto_do_arquivo("13825_DEBORA BRUDNA_17_09_2026_16_13_51.cutplanning") == "13825"
    assert projeto_do_arquivo("sem_numero.cutplanning") == "sem_numero"


def test_lote_com_varios_projetos_num_unico_plano(tmp_path):
    a, b, c, d = (_peca("13955", "GABRIELA", "5479 - SOLTEIRO"), _peca("13956", "GABRIELA", "5480 - BANHO SUITE"),
                  _peca("13957", "GABRIELA", "5481 - BANHO SOCIAL"), _peca("13958", "JOAO", "COZINHA"))
    arquivo = _plano(tmp_path, "13955_GABRIELA.cutplanning", {
        10: [_program(False, 2740, 1840, 2, (a, b)), _program(False, 2740, 1840, 1, (c,)),
             _program(True, 1500, 900, 1, (b, d))],
        20: [_program(True, 800, 600, 1, (c,))],
    })
    plano = ler_plano(arquivo)
    # o nome do arquivo só traz o 13955, mas o plano (lote) tem 4 projetos
    assert list(plano.projetos) == ["13955", "13956", "13957", "13958"]
    info = plano.projetos["13956"]
    assert (info.pedido, info.ambiente, info.cliente) == ("5480", "BANHO SUITE", "GABRIELA")
    assert (plano.projetos["13958"].pedido, plano.projetos["13958"].ambiente) == ("", "COZINHA")

    estoque = [_ret(1, 10, 1500, 900, "999"), _ret(2, 10, 1500, 900, "13958"), _ret(3, 20, 800, 600)]
    analise = analisar([plano], estoque, {10: Chapa(10, "15", "BRANCO TX"), 20: Chapa(20, "18", "CUMARU")})
    assert analise.projetos == ["13955", "13956", "13957", "13958"]
    assert analise.chapas[0].quantidade == 3 and analise.chapas[0].projetos == {"13955", "13956", "13957"}
    # reservado para qualquer projeto do lote tem prioridade sobre o reservado para outro
    compartilhado = next(r for r in analise.retalhos if r.codigo == 10)
    assert (compartilhado.projeto, compartilhado.retalho.id, compartilhado.situacao) == (
        "13956, 13958", 2, SITUACAO_PROJETO)
    assert [r.projeto for r in analise.retalhos] == ["13956, 13958", "13957"]  # ordem pelo 1º projeto


def test_plano_sem_etiquetas_usa_nome_do_arquivo(tmp_path):
    plano = ler_plano(_plano(tmp_path, "777_x.cutplanning", {10: [_program(False, 2740, 1840)]}))
    assert list(plano.projetos) == ["777"] and plano.materiais[0].programas[0].projetos == ["777"]


def test_chapa_em_posicao_diferente_e_a_mesma(tmp_path):
    a = _plano(tmp_path, "1_a.cutplanning", {10: [_program(False, 2740, 1840, 2), _program(False, 1840, 2740)]})
    analise = analisar([ler_plano(a)], [], {10: Chapa(10, "15", "BRANCO TX")})
    assert [(c.medida, c.quantidade) for c in analise.chapas] == [("2740 x 1840", 3)]


def test_material_sem_cadastro_nao_entra_nas_chapas(tmp_path):
    a = _plano(tmp_path, "1_a.cutplanning", {
        10: [_program(False, 2740, 1840, 2)],
        38015: [_program(False, 2740, 1840)],  # "MDF Retalho 15mm": não existe em tblChapas
    })
    analise = analisar([ler_plano(a)], [], {10: Chapa(10, "15", "BRANCO TX")})
    assert [(c.codigo, c.quantidade) for c in analise.chapas] == [(10, 2)]
    assert analise.total_chapas == 2
    assert [(c.codigo, c.quantidade) for c in analise.ignoradas] == [(38015, 1)]
    assert analise.texto_ignoradas == "MDF 38015 (1)"
    assert "Ignoradas (material sem cadastro): MDF 38015 (1)" in resumo_chapas(analise)
    wb = openpyxl.load_workbook(exportar_excel(analise, tmp_path / "c.xlsx"))
    assert any("MDF 38015" in str(c.value) for row in wb["Chapas inteiras"].iter_rows() for c in row)


def test_ordem_dos_retalhos(tmp_path):
    planos = [
        ler_plano(_plano(tmp_path, "100_c.cutplanning", {30: [_program(True, 500, 400)],
                                                         10: [_program(True, 600, 400)]})),
        ler_plano(_plano(tmp_path, "9_b.cutplanning", {20: [_program(True, 700, 400)],
                                                       10: [_program(True, 800, 400)]})),
    ]
    chapas = {10: Chapa(10, "15", "CARVALHO"), 20: Chapa(20, "15", "AMENDOA"), 30: Chapa(30, "6", "BRANCO")}
    analise = analisar(planos, [], chapas)
    assert [(r.projeto, r.descricao) for r in analise.retalhos] == [
        ("9", "AMENDOA"), ("9", "CARVALHO"), ("100", "BRANCO"), ("100", "CARVALHO"),
    ]


def test_analise(tmp_path):
    a = _plano(tmp_path, "100_cliente.cutplanning", {
        10: [_program(False, 2740, 1840, 2), _program(False, 2740, 1840),
             _program(True, 1500, 900), _program(True, 1500, 900)],
        20: [_program(True, 900, 1000), _program(True, 700, 700)],
    })
    b = _plano(tmp_path, "200_outro.cutplanning", {10: [_program(False, 2740, 1820)]})
    estoque = [
        _ret(1, 10, 1500, 900, "999"),   # reservado p/ outro projeto
        _ret(2, 10, 1500, 900, "100"),   # reservado p/ este projeto -> prioridade
        _ret(3, 20, 1000, 900),          # medidas invertidas, disponível
    ]
    analise = analisar([ler_plano(a), ler_plano(b)], estoque, {10: Chapa(10, "15", "BRANCO TX")})

    # chapas inteiras: uma linha por material + medida; material só com retalhos não aparece
    assert [(c.codigo, c.medida, c.quantidade, c.descricao) for c in analise.chapas] == [
        (10, "2740 x 1820", 1, "BRANCO TX"), (10, "2740 x 1840", 3, "BRANCO TX"),
    ]
    assert analise.chapas[1].projetos == {"100"}
    assert analise.total_chapas == 4

    # retalhos usados
    assert analise.total_retalhos == 4 and analise.nao_encontrados == 1
    situacoes = {(r.codigo, r.retalho.id if r.retalho else None): r.situacao for r in analise.retalhos}
    assert situacoes == {(10, 2): SITUACAO_PROJETO, (10, 1): SITUACAO_OUTRO,
                         (20, 3): SITUACAO_DISPONIVEL, (20, None): SITUACAO_NAO_ENCONTRADO}
    assert next(r for r in analise.retalhos if r.codigo == 20).descricao == "MDF 20"
    assert {r.setor for r in analise.retalhos if r.encontrado} == {"CORTE"}

    assert "3	10	BRANCO TX	15	2740 x 1840" in resumo_chapas(analise)
    assert "1500 x 900" in resumo_retalhos(analise) and "1500 x 900" not in resumo_chapas(analise)

    wb = openpyxl.load_workbook(exportar_excel(analise, tmp_path / "out" / "consumo.xlsx"))
    assert wb.sheetnames == ["Chapas inteiras", "Retalhos usados", "Arquivos"]
    assert list(wb["Chapas inteiras"].iter_rows(values_only=True))[-1][4] == 4
    ws = wb["Retalhos usados"]
    assert ws.max_row == 6  # cabeçalho + 4 + total
    assert [c.value for c in ws[1]] == ["Código", "Setor", "Esp. (mm)", "Material", "Medida", "Projeto",
                                        "ID no estoque", "Sobra do projeto"]
    falta = next(row for row in ws.iter_rows(min_row=2) if row[1].value == "NÃO ENCONTRADO NO ESTOQUE")
    assert falta[0].fill.fgColor.rgb.endswith("FDE2E2")
    assert "NÃO ENCONTRADO NO ESTOQUE" in resumo_retalhos(analise)
    assert "Situação" not in resumo_retalhos(analise)


def test_arquivo_real():
    exemplos = list(DADOS.glob("*.cutplanning"))
    if not exemplos:
        pytest.skip("sem arquivo de exemplo")
    plano = ler_plano(exemplos[0])
    assert plano.materiais
    cadastrados = {m.codigo: Chapa(m.codigo, fmt_number(m.espessura), m.descricao)
                   for m in plano.materiais if "RETALHO" not in m.descricao.upper()}
    analise = analisar([plano], [], cadastrados)
    assert analise.total_chapas > 0
    assert all("RETALHO" in c.descricao.upper() for c in analise.ignoradas)
    assert all(r.situacao == SITUACAO_NAO_ENCONTRADO for r in analise.retalhos)


def test_service_cruza_com_banco(settings, tmp_path):
    svc = EstoqueService(settings)
    rid = svc.lancar_retalho(10, 1500, 900)
    arquivo = _plano(tmp_path, "300_x.cutplanning", {10: [_program(True, 1500, 900), _program(False, 2740, 1840)]})
    analise = svc.analisar_planos([arquivo])
    assert (analise.chapas[0].descricao, analise.chapas[0].quantidade) == ("BRANCO TX", 1)
    assert analise.retalhos[0].retalho.id == rid and analise.retalhos[0].situacao == SITUACAO_DISPONIVEL
