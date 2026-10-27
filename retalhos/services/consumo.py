"""Consumo de material de planos de corte (``.cutplanning``) do Corte Certo.

Estrutura relevante do arquivo::

    cutplan
      material  id=<código>  description=...  thickness=...
        program  iswaste=false|true  lenght=...  width=...  quantity=<vezes>

- ``iswaste="false"``: o programa corta uma **chapa inteira** do estoque.
- ``iswaste="true"``: o programa corta um **retalho**; ``lenght`` × ``width`` são as medidas
  do retalho (as mesmas gravadas no ``RETxxxxx.TAB``). O arquivo não traz o ID do retalho,
  por isso ele é identificado no estoque por material + medidas.

O resultado é separado em duas listas: chapas inteiras (para compras) e retalhos usados.
"""
from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..models import Chapa, Retalho, Sobra, fmt_number
from .etiquetas import BARCODE_LENGTH, ler_codigo_barras

_P = "com.geeksystem.cutplanning.cutplan."
_TOLERANCIA = 0.05  # mm
_TOLERANCIA_SOBRA = 0.15  # mm: medida da etiqueta de sobra x medida do retalho no outro plano


# ================================================================== leitura
# Campos da etiqueta de peça (template 1) usados para identificar o projeto de cada peça
_CAMPO_IMAGEM = "172"   # projects/CLIENTE-images/\PLA<projeto>_<peça>.bmp
_CAMPO_CLIENTE = "184"
_CAMPO_AMBIENTE = "187"  # "<pedido> - <ambiente>"
_CAMPO_SOBRA = "160"     # etiqueta de sobra (template 2): código de barras de 21 caracteres
_RE_PROJETO_IMAGEM = re.compile(r"PLA(\d+)_", re.IGNORECASE)
_RE_PEDIDO_AMBIENTE = re.compile(r"^\s*(\d+)\s*-\s*(.+?)\s*$")


@dataclass
class InfoProjeto:
    """Dados de um projeto lidos das etiquetas das peças."""

    projeto: str
    pedido: str = ""
    ambiente: str = ""
    cliente: str = ""


@dataclass
class ProgramaCorte:
    comprimento: float
    largura: float
    quantidade: int
    retalho: bool
    projetos: list[str] = field(default_factory=list)  # projetos com peças neste programa


@dataclass
class MaterialPlano:
    codigo: int
    descricao: str
    espessura: float
    programas: list[ProgramaCorte] = field(default_factory=list)


@dataclass
class PlanoDeCorte:
    arquivo: Path
    projetos: dict[str, InfoProjeto]  # um plano (lote) pode ter vários projetos
    materiais: list[MaterialPlano]
    sobras: list[Sobra] = field(default_factory=list)  # sobras que este plano gera (etiquetas de sobra)

    @property
    def projeto(self) -> str:
        """Projetos do plano em texto (``"13825, 13826"``)."""
        return ", ".join(self.projetos)


def projeto_do_arquivo(arquivo: str | Path) -> str:
    """``13825_DEBORA BRUDNA_17_09_2026.cutplanning`` -> ``13825``."""
    nome = Path(arquivo).stem
    m = re.match(r"\s*(\d+)", nome)
    return m.group(1) if m else nome


def ordem_projeto(projeto: str) -> tuple[int, int, str]:
    """Projetos numéricos primeiro, em ordem numérica (9 antes de 10); depois os demais por nome."""
    return (0, int(projeto), "") if projeto.isdigit() else (1, 0, projeto.upper())


def _float(valor: str | None) -> float:
    try:
        return float(valor or 0)
    except ValueError:
        return 0.0


def _info_da_peca(data: ET.Element) -> InfoProjeto | None:
    campos = {f.get("name"): (f.get("value") or "").strip()
              for f in data.findall(_P + "material.program.cut.data.field")}
    m = _RE_PROJETO_IMAGEM.search(campos.get(_CAMPO_IMAGEM, ""))
    if not m:
        return None
    info = InfoProjeto(m.group(1), cliente=campos.get(_CAMPO_CLIENTE, ""))
    ambiente = campos.get(_CAMPO_AMBIENTE, "")
    pa = _RE_PEDIDO_AMBIENTE.match(ambiente)
    info.pedido, info.ambiente = (pa.group(1), pa.group(2)) if pa else ("", ambiente)
    return info


def ler_plano(arquivo: str | Path) -> PlanoDeCorte:
    root = ET.parse(str(arquivo)).getroot()
    projetos: dict[str, InfoProjeto] = {}
    materiais: list[MaterialPlano] = []
    sobras: list[Sobra] = []
    for mat in root.iter(_P + "material"):
        codigo = int(_float(mat.get("id")))
        mp = MaterialPlano(codigo, (mat.get("description") or "").strip(), _float(mat.get("thickness")))
        for prog in mat.findall(_P + "material.program"):
            do_programa: list[str] = []
            for data in prog.iter(_P + "material.program.cut.data"):
                if data.get("template") == "2":  # etiqueta de sobra gerada por este plano
                    sobra = _sobra_da_etiqueta(data)
                    if sobra:
                        sobras.append(sobra)
                    continue
                if data.get("template") != "1":
                    continue
                info = _info_da_peca(data)
                if info is None:
                    continue
                projetos.setdefault(info.projeto, info)
                if info.projeto not in do_programa:
                    do_programa.append(info.projeto)
            mp.programas.append(ProgramaCorte(
                comprimento=_float(prog.get("lenght")),  # grafia do Corte Certo
                largura=_float(prog.get("width")),
                quantidade=max(int(_float(prog.get("quantity"))), 1),
                retalho=(prog.get("iswaste") or "").lower() == "true",
                projetos=sorted(do_programa, key=ordem_projeto),
            ))
        materiais.append(mp)

    if not projetos:  # etiquetas sem a imagem da peça: usa o número no nome do arquivo
        num = projeto_do_arquivo(arquivo)
        projetos[num] = InfoProjeto(num)
    padrao = sorted(projetos, key=ordem_projeto)
    for mp in materiais:
        for prog in mp.programas:
            prog.projetos = prog.projetos or list(padrao)
    ordenados = {p: projetos[p] for p in padrao}
    return PlanoDeCorte(Path(arquivo), ordenados, materiais, sobras)


def _sobra_da_etiqueta(data: ET.Element) -> Sobra | None:
    for f in data.findall(_P + "material.program.cut.data.field"):
        valor = f.get("value") or ""
        if f.get("name") == _CAMPO_SOBRA and len(valor) == BARCODE_LENGTH:
            try:
                return ler_codigo_barras(valor)
            except ValueError:
                return None
    return None


# ================================================================== análise
SITUACAO_PROJETO = "Reservado p/ o projeto"
SITUACAO_DISPONIVEL = "Disponível no estoque"
SITUACAO_OUTRO = "Reservado p/ outro projeto"
SITUACAO_NAO_ENCONTRADO = "Não encontrado no estoque"


@dataclass
class ChapaInteira:
    """Quantidade de chapas inteiras de um material/medida (lista de compras)."""

    codigo: int
    descricao: str
    espessura: str
    comprimento: float
    largura: float
    quantidade: int = 0
    projetos: set[str] = field(default_factory=set)
    cadastrado: bool = True

    @property
    def medida(self) -> str:
        return f"{fmt_number(self.comprimento)} x {fmt_number(self.largura)}"


@dataclass
class RetalhoUsado:
    projeto: str
    codigo: int
    descricao: str
    espessura: str
    comprimento: float
    largura: float
    retalho: Retalho | None
    situacao: str
    origem: str = ""  # projeto(s) do lote cuja sobra este retalho é ("" = retalho do estoque)

    @property
    def encontrado(self) -> bool:
        return self.retalho is not None

    @property
    def passagem(self) -> tuple[str, str] | None:
        """``(origem, destino)`` quando o retalho é sobra de outro projeto do lote."""
        return (self.origem, self.projeto) if self.origem else None

    @property
    def setor(self) -> str:
        return self.retalho.setor if self.retalho else ""

    @property
    def medida(self) -> str:
        return f"{fmt_number(self.comprimento)} x {fmt_number(self.largura)}"


@dataclass
class AnaliseConsumo:
    planos: list[PlanoDeCorte]
    chapas: list[ChapaInteira]
    retalhos: list[RetalhoUsado]
    ignoradas: list[ChapaInteira] = field(default_factory=list)  # chapas de material sem cadastro

    @property
    def projetos(self) -> list[str]:
        """Todos os projetos dos planos (um lote pode ter vários), em ordem numérica."""
        return sorted({p for plano in self.planos for p in plano.projetos}, key=ordem_projeto)

    @property
    def info_projetos(self) -> dict[str, InfoProjeto]:
        info: dict[str, InfoProjeto] = {}
        for plano in self.planos:
            for num, dados in plano.projetos.items():
                info.setdefault(num, dados)
        return {p: info[p] for p in self.projetos}

    @property
    def total_chapas(self) -> int:
        return sum(c.quantidade for c in self.chapas)

    @property
    def total_retalhos(self) -> int:
        return len(self.retalhos)

    @property
    def nao_encontrados(self) -> int:
        return sum(not r.encontrado for r in self.retalhos)

    @property
    def texto_ignoradas(self) -> str:
        """``"MDF Retalho 15mm (1), MDF Retalho 6mm (1)"``."""
        return ", ".join(f"{c.descricao} ({c.quantidade})" for c in self.ignoradas)

    @property
    def passagens(self) -> list[tuple[str, str]]:
        """Sobras que passam de um projeto do lote para outro, em ordem (1º -> 2º, 2º -> 3º...)."""
        unicas = {r.passagem for r in self.retalhos if r.passagem}
        return sorted(unicas, key=lambda p: (ordem_projeto(_primeiro(p[0])), ordem_projeto(_primeiro(p[1])), p))

    def indice_passagem(self, r: RetalhoUsado) -> int | None:
        """Posição da passagem do retalho (0 = primeira cor, 1 = segunda...)."""
        return self.passagens.index(r.passagem) if r.passagem else None


def _primeiro(projetos: str) -> str:
    return projetos.split(",")[0].strip()


def _mesma_medida(r: Retalho, comprimento: float, largura: float) -> bool:
    def eq(a: float, b: float) -> bool:
        return abs(a - b) <= _TOLERANCIA
    return (eq(r.largura, comprimento) and eq(r.altura, largura)) or \
           (eq(r.largura, largura) and eq(r.altura, comprimento))


def _identificar(estoque: list[Retalho], usados: set[int], projetos: set[str], codigo: int,
                 comprimento: float, largura: float) -> tuple[Retalho | None, str]:
    candidatos = [r for r in estoque if r.codigo == codigo and r.id not in usados
                  and _mesma_medida(r, comprimento, largura)]
    for situacao, filtro in (
        (SITUACAO_PROJETO, lambda r: r.reserva in projetos),
        (SITUACAO_DISPONIVEL, lambda r: not r.reservado),
        (SITUACAO_OUTRO, lambda r: True),
    ):
        escolhidos = sorted((r for r in candidatos if filtro(r)), key=lambda r: r.id)
        if escolhidos:
            return escolhidos[0], situacao
    return None, SITUACAO_NAO_ENCONTRADO


def _mesmas_medidas(a1: float, b1: float, a2: float, b2: float) -> bool:
    def eq(x: float, y: float) -> bool:
        return abs(x - y) <= _TOLERANCIA_SOBRA
    return (eq(a1, a2) and eq(b1, b2)) or (eq(a1, b2) and eq(b1, a2))


def _origem_da_sobra(sobras: list[tuple[int, str, Sobra]], usadas: set[int], indice_plano: int,
                     codigo: int, comprimento: float, largura: float) -> str:
    """Projeto(s) de outro plano do lote que geraram uma sobra igual a este retalho ("" se nenhum)."""
    for n, (plano, origem, sobra) in enumerate(sobras):
        if n in usadas or plano == indice_plano or sobra.codigo != codigo:
            continue
        if _mesmas_medidas(sobra.largura, sobra.altura, comprimento, largura):
            usadas.add(n)  # cada sobra só vira um retalho
            return origem
    return ""


def analisar(planos: list[PlanoDeCorte], estoque: list[Retalho], chapas: dict[int, Chapa]) -> AnaliseConsumo:
    """Separa os programas em chapas inteiras (somadas por material/medida) e retalhos usados.

    Retalho usado por um plano que bate (material + medida) com uma sobra gerada por **outro**
    plano carregado junto é marcado como sobra daquele projeto (``origem``): é o caso de lotes
    em que a sobra do 1º projeto é usada no 2º, a do 2º no 3º, e assim por diante.
    Retalho encontrado no estoque **com setor** nunca é considerado sobra (é retalho guardado);
    sobras lançadas pelo Leitor XML entram sem setor.
    """
    inteiras: dict[tuple[int, float, float], ChapaInteira] = {}
    retalhos: list[RetalhoUsado] = []
    usados: set[int] = set()
    sobras = [(i, plano.projeto, s) for i, plano in enumerate(planos) for s in plano.sobras]
    sobras_usadas: set[int] = set()

    for indice_plano, plano in enumerate(planos):
        do_lote = set(plano.projetos)  # retalho reservado para qualquer projeto do lote tem prioridade
        for mp in plano.materiais:
            cad = chapas.get(mp.codigo)
            descricao = cad.descricao if cad else mp.descricao
            espessura = cad.espessura if cad else fmt_number(mp.espessura)
            for prog in mp.programas:
                if not prog.retalho:
                    # 2740 x 1840 e 1840 x 2740 são a mesma chapa: normaliza para maior x menor
                    maior, menor = max(prog.comprimento, prog.largura), min(prog.comprimento, prog.largura)
                    chave = (mp.codigo, maior, menor)
                    ci = inteiras.get(chave)
                    if ci is None:
                        ci = inteiras[chave] = ChapaInteira(
                            mp.codigo, descricao, espessura, maior, menor,
                            cadastrado=cad is not None,
                        )
                    ci.quantidade += prog.quantidade
                    ci.projetos.update(prog.projetos)
                    continue
                for _ in range(prog.quantidade):
                    r, situacao = _identificar(estoque, usados, do_lote, mp.codigo,
                                               prog.comprimento, prog.largura)
                    if r:
                        usados.add(r.id)
                    # retalho do estoque com setor é físico/guardado: nunca é sobra de outro projeto, mesmo
                    # com a mesma medida (e não consome a sobra, que fica para outra ocorrência igual)
                    origem = "" if r is not None and r.setor else _origem_da_sobra(
                        sobras, sobras_usadas, indice_plano, mp.codigo, prog.comprimento, prog.largura)
                    retalhos.append(RetalhoUsado(", ".join(prog.projetos), mp.codigo, descricao, espessura,
                                                 prog.comprimento, prog.largura, r, situacao, origem))

    ordem = sorted(inteiras.values(), key=lambda c: (c.descricao, _ordem_espessura(c.espessura), c.codigo, c.medida))
    # material sem cadastro em tblChapas (ex.: "MDF Retalho 6mm") não entra na lista de compras
    lista_chapas = [c for c in ordem if c.cadastrado]
    ignoradas = [c for c in ordem if not c.cadastrado]
    # projeto em ordem numérica crescente, depois material em ordem alfabética
    # retalho compartilhado ("13825, 13826") ordena pelo primeiro projeto
    retalhos.sort(key=lambda r: (ordem_projeto(r.projeto.split(",")[0].strip()), r.projeto,
                                 r.descricao.upper(), r.codigo, -r.comprimento, -r.largura))
    return AnaliseConsumo(planos, lista_chapas, retalhos, ignoradas)


def _ordem_espessura(espessura: str) -> float:
    try:
        return float(espessura)
    except ValueError:
        return float("inf")


# ================================================================ exportação
def resumo_chapas(analise: AnaliseConsumo) -> str:
    """Lista de chapas inteiras em texto (tabulado) para colar em e-mail/WhatsApp/planilha."""
    linhas = [f"Chapas inteiras — projeto(s): {', '.join(analise.projetos)}", "",
              "Qtd\tCódigo\tMaterial\tEsp.\tMedida"]
    for c in analise.chapas:
        linhas.append(f"{c.quantidade}\t{c.codigo}\t{c.descricao}\t{c.espessura}\t{c.medida}")
    linhas.append(f"{analise.total_chapas}\tTOTAL")
    if analise.ignoradas:
        linhas += ["", f"Ignoradas (material sem cadastro): {analise.texto_ignoradas}"]
    return "\n".join(linhas)


def resumo_retalhos(analise: AnaliseConsumo) -> str:
    """Lista de retalhos usados em texto (tabulado)."""
    linhas = [f"Retalhos usados — projeto(s): {', '.join(analise.projetos)}", "",
              "Código\tSetor\tEsp.\tMaterial\tMedida\tProjeto\tID estoque"]
    for r in analise.retalhos:
        rid, setor = (r.retalho.id, r.setor) if r.retalho else ("-", "NÃO ENCONTRADO NO ESTOQUE")
        linhas.append(f"{r.codigo}\t{setor}\t{r.espessura}\t{r.descricao}\t{r.medida}\t{r.projeto}\t{rid}")
    linhas.append(f"TOTAL\t{analise.total_retalhos}")
    return "\n".join(linhas)


def exportar_excel(analise: AnaliseConsumo, destino: str | Path) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    negrito = Font(bold=True)
    cabecalho = PatternFill("solid", fgColor="DBE7FF")
    nao_encontrado = PatternFill("solid", fgColor="FDE2E2")

    def preparar(ws, titulos: list[str], larguras: list[int]) -> None:
        ws.append(titulos)
        for cell, w in zip(ws[1], larguras, strict=True):
            cell.font = negrito
            cell.fill = cabecalho
            ws.column_dimensions[cell.column_letter].width = w
        ws.freeze_panes = "A2"

    def total(ws, valores: list) -> None:
        ws.append(valores)
        for cell in ws[ws.max_row]:
            cell.font = negrito

    wb = Workbook()
    ws = wb.active
    ws.title = "Chapas inteiras"
    preparar(ws, ["Código", "Material", "Esp. (mm)", "Medida da chapa", "Quantidade", "Projetos"],
             [10, 38, 10, 18, 12, 24])
    for c in analise.chapas:
        ws.append([c.codigo, c.descricao, c.espessura, c.medida, c.quantidade, ", ".join(sorted(c.projetos))])
    total(ws, ["", "TOTAL", "", "", analise.total_chapas, ""])
    if analise.ignoradas:
        ws.append([])
        ws.append(["", f"Ignoradas (material sem cadastro): {analise.texto_ignoradas}"])
        ws.cell(ws.max_row, 2).font = Font(italic=True, color="6B7280")

    ws = wb.create_sheet("Retalhos usados")
    preparar(ws, ["Código", "Setor", "Esp. (mm)", "Material", "Medida", "Projeto", "ID no estoque",
                  "Sobra do projeto"], [10, 26, 10, 38, 16, 10, 14, 16])
    for r in analise.retalhos:
        ws.append([r.codigo, (r.setor or None) if r.retalho else "NÃO ENCONTRADO NO ESTOQUE",
                   r.espessura, r.descricao, r.medida, r.projeto, r.retalho.id if r.retalho else None,
                   r.origem or None])
        if not r.encontrado and not r.origem:
            for cell in ws[ws.max_row]:
                cell.fill = nao_encontrado
    total(ws, ["TOTAL", analise.total_retalhos, "", "", "", "", "", ""])

    ws = wb.create_sheet("Arquivos")
    preparar(ws, ["Projeto", "Arquivo"], [10, 90])
    for p in analise.planos:
        ws.append([p.projeto, str(p.arquivo)])

    wb.save(destino)
    return destino
