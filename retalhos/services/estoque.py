"""Regras de negócio do estoque: retalhos, materiais, reservas e baixas.

Toda operação que altera um retalho mantém o banco e o arquivo ``RETxxxxx.TAB``
sincronizados: o Corte Certo só deve oferecer retalhos disponíveis.
"""
from __future__ import annotations

import shutil
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..config import DATE_FORMAT, Settings
from ..db import ChapaRepository, Database, HistoricoRepository, RetalhoRepository
from ..models import Chapa, Movimento, Retalho, RetalhoRemovido, Sobra
from .capa_lote import CapaLote, capa_de_analise, ler_conf_folha_lote
from .consumo import AnaliseConsumo, RetalhoUsado, analisar, ler_plano
from .corte_certo import ArquivosCorteCerto


# ações gravadas no histórico (tblHistorico)
ACAO_LANCAMENTO = "LANÇAMENTO"
ACAO_SOBRA = "SOBRA LANÇADA"
ACAO_RESERVA = "RESERVA"
ACAO_LIBERACAO = "LIBERAÇÃO"
ACAO_CANCELAMENTO = "CANCELAMENTO"
ACAO_BAIXA = "BAIXA"
ACAO_EXCLUSAO = "EXCLUSÃO"
ACAO_RESTAURACAO = "RESTAURAÇÃO"


class RegraDeNegocioError(Exception):
    """Operação recusada por regra de negócio (mensagem pronta para o usuário)."""


@dataclass
class ResultadoReserva:
    reservados: int
    planilha: Path | None


@dataclass
class ResultadoReservaPlano:
    reservados: list[tuple[int, str]] = field(default_factory=list)  # (id do retalho, projeto)
    ja_reservados: int = 0
    conflitos: list[str] = field(default_factory=list)
    nao_encontrados: int = 0


def projeto_da_reserva(item: RetalhoUsado) -> str:
    """Retalho usado por mais de um projeto do lote fica com o primeiro (``"13956, 13958"`` -> ``13956``)."""
    return item.projeto.split(",")[0].strip()


class EstoqueService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.db = Database(settings.database_path, settings.database_password)
        self.arquivos = ArquivosCorteCerto(settings)

    def _agora(self) -> datetime:
        return datetime.now().replace(microsecond=0)

    # ============================================================ materiais
    def listar_materiais(self) -> list[Chapa]:
        with self.db.session() as s:
            return ChapaRepository(s).listar()

    def buscar_material(self, codigo: int) -> Chapa | None:
        with self.db.session() as s:
            return ChapaRepository(s).buscar(codigo)

    def buscar_materiais(self, codigos: Iterable[int]) -> dict[int, Chapa]:
        with self.db.session() as s:
            return ChapaRepository(s).buscar_varios(codigos)

    def salvar_material(self, chapa: Chapa) -> bool:
        """Insere ou atualiza. Retorna ``True`` se foi um novo cadastro."""
        if not chapa.descricao.strip():
            raise RegraDeNegocioError("Informe a descrição do material.")
        with self.db.session() as s:
            repo = ChapaRepository(s)
            if repo.buscar(chapa.codigo) is None:
                repo.inserir(chapa)
                return True
            repo.atualizar(chapa)
            return False

    # ============================================================= retalhos
    def listar_retalhos(self) -> list[Retalho]:
        with self.db.session() as s:
            return RetalhoRepository(s).listar()

    def validar_medidas(self, largura: float, altura: float) -> None:
        st = self.settings
        if not st.min_width <= largura <= st.max_width:
            raise RegraDeNegocioError(
                f"Largura fora do limite permitido (mín. {st.min_width:g} e máx. {st.max_width:g})."
            )
        if not st.min_height <= altura <= st.max_height:
            raise RegraDeNegocioError(
                f"Altura fora do limite permitido (mín. {st.min_height:g} e máx. {st.max_height:g})."
            )

    def lancar_retalho(self, codigo: int, largura: float, altura: float, setor: str = "") -> int:
        """Registra um retalho novo e o disponibiliza para o Corte Certo."""
        self.validar_medidas(largura, altura)
        self.arquivos.garantir_pasta_retalhos()
        with self.db.session() as s:
            if ChapaRepository(s).buscar(codigo) is None:
                raise RegraDeNegocioError(f"Material {codigo} não cadastrado.")
            agora = self._agora()
            novo_id = RetalhoRepository(s).inserir(
                codigo=codigo, largura=largura, altura=altura, setor=setor.strip() or None,
                data_entrada=agora, responsavel=self.settings.user,
            )
            self._historico(s, novo_id, ACAO_LANCAMENTO, f"{largura:g} x {altura:g}"
                            + (f", setor {setor.strip()}" if setor.strip() else ""), agora)
            # dentro da transação: se o arquivo falhar, o registro é desfeito
            self.arquivos.adicionar(codigo, novo_id, largura, altura)
        return novo_id

    # ============================================================ histórico
    def _historico(self, s, retalho_id: int, acao: str, detalhe: str = "", quando: datetime | None = None) -> None:
        """Grava a ação na mesma transação da alteração (as duas valem ou nenhuma)."""
        HistoricoRepository(s).registrar(retalho_id, acao, quando or self._agora(), self.settings.user, detalhe)

    def historico(self, retalho_id: int) -> list[Movimento]:
        with self.db.session() as s:
            return HistoricoRepository(s).listar(retalho_id)

    def lancar_sobras(self, sobras: Iterable[Sobra], projeto: str) -> int:
        """Registra sobras de um plano de corte já reservadas para ``projeto``."""
        projeto = projeto.strip()
        if not projeto:
            raise RegraDeNegocioError("Informe o projeto para o qual as sobras ficarão reservadas.")
        self.arquivos.garantir_pasta_retalhos()
        agora = self._agora()
        total = 0
        with self.db.session() as s:
            repo = RetalhoRepository(s)
            for sobra in sobras:
                novo_id = repo.inserir(
                    codigo=sobra.codigo, largura=sobra.largura, altura=sobra.altura,
                    data_entrada=agora, responsavel=self.settings.user,
                    reserva=projeto, data_reserva=agora.strftime(DATE_FORMAT),
                )
                self._historico(s, novo_id, ACAO_SOBRA,
                                f"{sobra.largura:g} x {sobra.altura:g}, reservada p/ {projeto}", agora)
                self.arquivos.adicionar(sobra.codigo, novo_id, sobra.largura, sobra.altura, projeto)
                total += 1
        return total

    def excluir_retalho(self, retalho: Retalho) -> None:
        """Exclusão registrada: não apaga; grava data/hora e quem excluiu e o retalho some do sistema."""
        if retalho.reservado:
            raise RegraDeNegocioError("Não é possível excluir um retalho reservado.")
        if not self.settings.user:
            raise RegraDeNegocioError("É preciso estar logado para excluir.")
        self.arquivos.garantir_pasta_retalhos()
        agora = self._agora()
        with self.db.session() as s:
            RetalhoRepository(s).excluir(retalho.id, agora, self.settings.user)
            self._historico(s, retalho.id, ACAO_EXCLUSAO,
                            f"{retalho.descricao} {retalho.largura:g} x {retalho.altura:g}", agora)
            self.arquivos.remover(retalho.codigo, retalho.id)

    def liberar_reserva(self, retalho: Retalho) -> None:
        """Devolve um retalho reservado ao estoque disponível.

        Só limpa a reserva: a data de lançamento (DataEntrada) e o responsável continuam os originais.
        """
        if not self.settings.allow_remove_reserve:
            raise RegraDeNegocioError("Você não tem permissão para remover reservas.")
        if not retalho.reservado:
            raise RegraDeNegocioError("Este retalho não está reservado.")
        self.arquivos.garantir_pasta_retalhos()
        with self.db.session() as s:
            RetalhoRepository(s).cancelar_reserva(retalho.id)
            self._historico(s, retalho.id, ACAO_LIBERACAO, self._texto_reserva(retalho))
            self.arquivos.adicionar(retalho.codigo, retalho.id, retalho.largura, retalho.altura)

    @staticmethod
    def _texto_reserva(r: Retalho) -> str:
        """``"estava reservado p/ 1206 (por MARIA)"``: o histórico guarda de quem era a reserva."""
        return f"estava reservado p/ {r.reserva}" + (f" (por {r.reservado_por})" if r.reservado_por else "")

    def recriar_arquivos(self) -> int:
        """Recria todos os ``RET*.TAB`` com os retalhos disponíveis do banco (somente administrador)."""
        if not self.settings.is_admin:
            raise RegraDeNegocioError("Somente administradores podem recriar os arquivos do Corte Certo.")
        with self.db.session() as s:
            disponiveis = RetalhoRepository(s).listar_disponiveis()
        return self.arquivos.recriar(disponiveis)

    # ============================================================= reservas
    def retalhos_para_projeto(self, projeto: str) -> list[Retalho]:
        """Retalhos dos materiais do pedido que estão livres ou já reservados para o projeto."""
        codigos = self.arquivos.codigos_do_pedido(projeto)
        if not codigos:
            return []
        with self.db.session() as s:
            retalhos = RetalhoRepository(s).listar_por_materiais(codigos)
        return [r for r in retalhos if r.reserva in ("", projeto)]

    def corte_gerado(self, projeto: str) -> bool:
        """Existe pasta de lote em "Projeto gerado" começando com o número do projeto?"""
        pasta = Path(self.settings.projects_dir)
        try:
            return any(item.name.startswith(projeto) for item in pasta.iterdir())
        except OSError:
            return False

    def reservar(self, projeto: str, retalhos: list[Retalho], gerar_planilha: bool = True,
                 origem: str = "pedido") -> ResultadoReserva:
        """Reserva os retalhos para ``projeto`` (texto livre: pode ser um projeto que ainda será feito).

        ``gerar_planilha=False`` é usado na reserva manual: a ``PRJ_<projeto>.xlsx`` seria
        sobrescrita só com estes retalhos.
        """
        projeto = projeto.strip().upper()
        if not projeto:
            raise RegraDeNegocioError("Informe o número do projeto.")
        if len(projeto) > 50:
            raise RegraDeNegocioError("O número/identificação do projeto deve ter no máximo 50 caracteres.")
        if not retalhos:
            raise RegraDeNegocioError("Selecione ao menos um retalho.")
        outros = [r for r in retalhos if r.reserva not in ("", projeto)]
        if outros:
            raise RegraDeNegocioError(
                f"O retalho {outros[0].id} já está reservado para o projeto {outros[0].reserva}."
            )
        self.arquivos.garantir_pasta_retalhos()
        data = self._agora().strftime(DATE_FORMAT)
        with self.db.session() as s:
            repo = RetalhoRepository(s)
            # 1º confere todos no banco (outra máquina pode ter reservado/baixado depois da tela carregar);
            # só depois grava, para não tirar linhas do RET de uma reserva que vai ser desfeita
            for r in retalhos:
                atual = repo.buscar(r.id)
                if atual is None:
                    raise RegraDeNegocioError(f"O retalho {r.id} não está mais no estoque. Atualize a lista.")
                if atual.reserva not in ("", projeto):
                    raise RegraDeNegocioError(
                        f"O retalho {r.id} acabou de ser reservado para o projeto {atual.reserva}. Atualize a lista."
                    )
            for r in retalhos:
                repo.reservar(r.id, projeto, data, self.settings.user)
                self._historico(s, r.id, ACAO_RESERVA, f"p/ {projeto} ({origem})")
                self.arquivos.remover(r.codigo, r.id)
        planilha = self._gerar_planilha(projeto, retalhos) if gerar_planilha else None
        return ResultadoReserva(len(retalhos), planilha)

    def reservar_manual(self, projeto: str, retalhos: list[Retalho]) -> ResultadoReserva:
        """Reserva livre (tela principal), para um projeto que pode ainda nem existir.

        Exige ``ReservarManual=Ativado`` no Settings.json. Não gera a ``PRJ_<projeto>.xlsx``.
        """
        if not self.settings.allow_manual_reserve:
            raise RegraDeNegocioError("Você não tem permissão para fazer reserva manual.")
        return self.reservar(projeto, retalhos, gerar_planilha=False, origem="manual")

    def cancelar_reservas(self, projeto: str, retalhos: list[Retalho]) -> int:
        """Desfaz reservas do projeto e devolve os retalhos ao arquivo do Corte Certo."""
        alvo = [r for r in retalhos if r.reserva == projeto]
        if not alvo:
            raise RegraDeNegocioError("Nenhum dos retalhos selecionados está reservado para este projeto.")
        self.arquivos.garantir_pasta_retalhos()
        with self.db.session() as s:
            repo = RetalhoRepository(s)
            for r in alvo:
                repo.cancelar_reserva(r.id)
                self._historico(s, r.id, ACAO_CANCELAMENTO, self._texto_reserva(r))
                self.arquivos.adicionar(r.codigo, r.id, r.largura, r.altura)
        return len(alvo)

    def retalhos_reservados(self, projeto: str) -> list[Retalho]:
        with self.db.session() as s:
            return RetalhoRepository(s).listar_por_reserva(projeto.strip())

    def dar_baixa(self, projeto: str) -> int:
        """Baixa por consumo dos retalhos reservados para o projeto.

        Não apaga: registra data/hora e responsável (usuário logado) e o retalho some de todo o
        sistema (listas, reservas, consumo, capa e arquivos RET). O registro fica no banco.
        """
        projeto = projeto.strip()
        if not self.settings.user:
            raise RegraDeNegocioError("É preciso estar logado para dar baixa.")
        self.arquivos.garantir_pasta_retalhos()
        agora = self._agora()
        with self.db.session() as s:
            repo = RetalhoRepository(s)
            retalhos = repo.listar_por_reserva(projeto)
            for r in retalhos:
                repo.dar_baixa(r.id, agora, self.settings.user)
                self._historico(s, r.id, ACAO_BAIXA, f"consumido no projeto {projeto}", agora)
                self.arquivos.remover(r.codigo, r.id)
        return len(retalhos)

    # ============================================================ removidos
    def contar_removidos_por_material(self) -> dict[int, int]:
        with self.db.session() as s:
            return RetalhoRepository(s).contar_removidos_por_material()

    def listar_removidos(self, codigo: int) -> list[RetalhoRemovido]:
        """Retalhos que receberam baixa ou foram excluídos (consulta sempre por material)."""
        with self.db.session() as s:
            return RetalhoRepository(s).listar_removidos(codigo)

    def restaurar_retalho(self, removido: RetalhoRemovido) -> None:
        """Devolve ao estoque, disponível, um retalho baixado/excluído por engano (somente administrador)."""
        if not self.settings.is_admin:
            raise RegraDeNegocioError("Somente administradores podem restaurar retalhos.")
        self.arquivos.garantir_pasta_retalhos()
        r = removido.retalho
        with self.db.session() as s:
            repo = RetalhoRepository(s)
            atual = repo.buscar_removido(r.id)
            if atual is None:
                raise RegraDeNegocioError(f"O retalho {r.id} já voltou ao estoque. Atualize a lista.")
            repo.restaurar(r.id)
            if atual.baixa:
                detalhe = f"estava baixado no projeto {atual.retalho.reserva}" if atual.retalho.reserva \
                    else "estava baixado"
            else:
                detalhe = "estava excluído"
            data = atual.data.strftime("%d/%m/%Y %H:%M") if atual.data else "—"
            self._historico(s, r.id, ACAO_RESTAURACAO, f"{detalhe} ({data} por {atual.responsavel or '—'})")
            self.arquivos.adicionar(r.codigo, r.id, r.largura, r.altura)

    @staticmethod
    def exportar_removidos(removidos: list[RetalhoRemovido], destino: str | Path) -> Path:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Retalhos removidos"
        titulos = ["ID", "Código", "Material", "Esp. (mm)", "Largura", "Altura", "Setor", "Motivo", "Data/hora",
                   "Responsável", "Projeto", "Lançado em", "Lançado por"]
        ws.append(titulos)
        for cell, w in zip(ws[1], [8, 10, 34, 10, 10, 10, 10, 10, 18, 16, 14, 18, 16], strict=True):
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DBE7FF")
            ws.column_dimensions[cell.column_letter].width = w
        ws.freeze_panes = "A2"
        for x in removidos:
            r = x.retalho
            ws.append([r.id, r.codigo, r.descricao, r.espessura, r.largura, r.altura, r.setor, x.motivo, x.data,
                       x.responsavel, r.reserva, r.data_entrada, r.responsavel])
        for linha in ws.iter_rows(min_row=2):
            for cell in (linha[8], linha[11]):
                cell.number_format = "dd/mm/yyyy hh:mm"
        wb.save(destino)
        return destino

    # ============================================================ migração
    def preparar_banco(self, pasta_backup: Path | None = None) -> Path | None:
        """Cria os controles que faltarem (1ª abertura da versão nova).

        - colunas em tblRetalhos: ResponsavelReserva, DataBaixa/ResponsavelBaixa, DataExclusao/ResponsavelExclusao;
        - tabela tblHistorico.
        Só acrescenta (nada existente é alterado). Antes de mexer num banco Access, guarda uma cópia em
        ``%LOCALAPPDATA%\\Retalhos\\backups``. Retorna o caminho do backup quando houve migração.
        """
        with self.db.session() as s:
            if not RetalhoRepository(s).precisa_migrar() and not HistoricoRepository(s).precisa_criar():
                return None
        backup = None
        if self.db.dialect == "access":
            from ..paths import data_dir
            pasta = pasta_backup or data_dir() / "backups"
            pasta.mkdir(parents=True, exist_ok=True)
            original = Path(self.settings.database_path)
            backup = pasta / f"{original.stem}_antes_controles_{self._agora():%Y%m%d_%H%M%S}{original.suffix}"
            shutil.copy2(original, backup)
        with self.db.session() as s:
            RetalhoRepository(s).migrar()
            HistoricoRepository(s).garantir_tabela()
        return backup

    # ============================================================== consumo
    def analisar_planos(self, arquivos: Iterable[str | Path]) -> AnaliseConsumo:
        """Chapas inteiras por material e retalhos usados em um ou mais ``.cutplanning``."""
        planos = [ler_plano(a) for a in arquivos]
        codigos = {m.codigo for p in planos for m in p.materiais}
        with self.db.session() as s:
            estoque = RetalhoRepository(s).listar_por_materiais(codigos)
            chapas = ChapaRepository(s).buscar_varios(codigos)
        return analisar(planos, estoque, chapas)

    def reservar_do_plano(self, analise: AnaliseConsumo,
                          itens: Iterable[RetalhoUsado] | None = None) -> ResultadoReservaPlano:
        """Reserva os retalhos que o plano de corte vai usar (identificados na análise).

        - cada retalho fica reservado para um único projeto: o primeiro que o usa no plano;
        - medidas repetidas já vêm casadas com IDs diferentes pela análise (um retalho por ocorrência);
        - o estado é conferido de novo no banco: retalho reservado por outra máquina depois da
          análise não é tomado (vira conflito); se já está reservado para um projeto do lote, nada muda.
        """
        resultado = ResultadoReservaPlano()
        do_lote = set(analise.projetos)
        alvo = []
        for item in analise.retalhos if itens is None else itens:
            if item.retalho is None:
                resultado.nao_encontrados += 1
            else:
                alvo.append(item)
        if not alvo:
            return resultado

        self.arquivos.garantir_pasta_retalhos()
        data = self._agora().strftime(DATE_FORMAT)
        with self.db.session() as s:
            repo = RetalhoRepository(s)
            for item in alvo:
                rid = item.retalho.id  # type: ignore[union-attr]
                atual = repo.buscar(rid)
                if atual is None:
                    resultado.conflitos.append(f"Retalho {rid} ({item.medida}) não está mais no estoque.")
                    continue
                if atual.reserva in do_lote:
                    resultado.ja_reservados += 1
                    continue
                if atual.reserva:
                    resultado.conflitos.append(
                        f"Retalho {rid} ({atual.descricao} {item.medida}) está reservado para o projeto {atual.reserva}."
                    )
                    continue
                projeto = projeto_da_reserva(item)
                repo.reservar(rid, projeto, data, self.settings.user)
                self._historico(s, rid, ACAO_RESERVA, f"p/ {projeto} (plano de corte)")
                self.arquivos.remover(atual.codigo, rid)  # sai do RET: outra otimização não o escolhe
                resultado.reservados.append((rid, projeto))
        return resultado

    def capa_do_lote(self, analise: AnaliseConsumo, lote: str = "", observacoes: str = "") -> CapaLote:
        """Capa do lote a partir da análise; projetos completados com o ``ConFolhaLote.xlsx``.

        O número do lote só é sugerido pelo ``ConFolhaLote`` se algum projeto do plano estiver nele.
        """
        lote_conf, projetos_conf = "", []
        try:
            lote_conf, projetos_conf = ler_conf_folha_lote(self.settings.conf_folha_lote)
        except Exception:  # noqa: BLE001 - planilha opcional (rede fora, arquivo aberto/inválido)
            pass
        if not lote and {p.projeto for p in projetos_conf} & set(analise.projetos):
            lote = lote_conf
        return capa_de_analise(analise, projetos_conf, lote, observacoes)

    def _gerar_planilha(self, projeto: str, retalhos: list[Retalho]) -> Path | None:
        """``PRJ_<projeto>.xlsx`` no mesmo layout de colunas da versão anterior."""
        if not self.settings.sheets_dir:
            return None
        pasta = Path(self.settings.sheets_dir)
        pasta.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        for r in retalhos:
            entrada = r.data_entrada.strftime(DATE_FORMAT) if r.data_entrada else None
            ws.append([
                r.setor or None, r.espessura, r.descricao, None, None, None, None,
                str(r.largura), "x", str(r.altura), projeto, r.responsavel or None, entrada,
            ])
        destino = pasta / f"PRJ_{projeto}.xlsx"
        wb.save(destino)
        return destino
