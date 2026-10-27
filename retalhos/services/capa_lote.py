"""Capa do lote de produção (novo layout, ``.xlsx`` sem macros, pronto para imprimir em A4).

Substitui o ``LOTE.xlsm``. Tudo vem do que o sistema gera:

- chapas inteiras e retalhos: análise do plano de corte (:mod:`.consumo`);
- número do lote e dados dos projetos: ``ConFolhaLote.xlsx`` (:func:`ler_conf_folha_lote`).

Impressão frente e verso: a pasta tem duas abas, ``Frente`` e ``Verso`` (controle de montagem),
e cada uma é ajustada para caber em exatamente 1 página A4. Assim, por mais itens que o lote
tenha, a montagem nunca "escorrega" para fora do verso. Quando o conteúdo não cabe no tamanho
normal, a face é refeita no modo compacto; só depois disso o Excel reduz a escala.
"""
from __future__ import annotations

import os
import subprocess
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from ..models import fmt_number
from .consumo import AnaliseConsumo


# ==================================================================== dados
@dataclass
class ProjetoLote:
    projeto: str
    pedido: str = ""
    ambiente: str = ""
    cliente: str = ""
    entrega: str = ""
    no_conf_folha: bool = True  # False: dados vieram só das etiquetas do plano de corte


@dataclass
class ChapaLote:
    material: str
    espessura: str
    quantidade: int
    medida: str = ""


@dataclass
class RetalhoLote:
    setor: str
    espessura: str
    material: str
    medida: str
    projeto: str
    encontrado: bool = True  # False: o plano usa um retalho que não está no estoque
    passagem: int | None = None  # índice em CapaLote.passagens: sobra de outro projeto do lote


@dataclass
class CapaLote:
    lote: str
    emitido: datetime
    projetos: list[ProjetoLote] = field(default_factory=list)
    chapas: list[ChapaLote] = field(default_factory=list)
    retalhos: list[RetalhoLote] = field(default_factory=list)
    observacoes: str = ""
    passagens: list[tuple[str, str]] = field(default_factory=list)  # (origem, destino) das sobras no lote

    @property
    def total_chapas(self) -> int:
        return sum(c.quantidade for c in self.chapas)

    @property
    def total_retalhos(self) -> int:
        return len(self.retalhos)

    @property
    def proxima_entrega(self) -> str:
        datas = []
        for p in self.projetos:
            try:
                datas.append(datetime.strptime(p.entrega, "%d/%m/%Y"))
            except ValueError:
                continue
        return min(datas).strftime("%d/%m/%Y") if datas else ""


# ============================================================ fontes de dados
def _texto(valor: Any) -> str:
    if valor is None or valor == 0 or isinstance(valor, time):
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float):
        return fmt_number(valor)
    return str(valor).strip()


def ler_conf_folha_lote(arquivo: str | Path) -> tuple[str, list[ProjetoLote]]:
    """Número do lote e projetos a partir do ``ConFolhaLote.xlsx`` (C.C., Pedido, Ambiente, Cliente, Entrega, Lote)."""
    wb = load_workbook(str(arquivo), data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        lote, projetos = "", []
        for linha in ws.iter_rows(min_row=2, values_only=True):
            if not linha or not _texto(linha[0]):
                continue
            projetos.append(ProjetoLote(*(_texto(v) for v in linha[:5])))
            lote = lote or (_texto(linha[5]) if len(linha) > 5 else "")
    finally:
        wb.close()
    return lote, projetos


def capa_de_analise(analise: AnaliseConsumo, conf_folha: list[ProjetoLote] | None = None,
                    lote: str = "", observacoes: str = "") -> CapaLote:
    """Monta a capa com o que o sistema gerou: chapas inteiras e retalhos do plano de corte.

    Os projetos são todos os encontrados no(s) plano(s) - um lote pode ter vários. Os dados
    (pedido, ambiente, cliente, entrega) vêm do ``ConFolhaLote.xlsx`` quando o projeto está lá;
    senão, das etiquetas das peças (sem a data de entrega). Cada retalho é único: uma linha por retalho.
    """
    conf = {p.projeto: p for p in conf_folha or []}
    projetos = []
    for num, info in analise.info_projetos.items():
        if num in conf:
            projetos.append(conf[num])
        else:
            projetos.append(ProjetoLote(num, info.pedido, info.ambiente, info.cliente, no_conf_folha=False))
    chapas = [ChapaLote(c.descricao, c.espessura, c.quantidade, c.medida) for c in analise.chapas]
    retalhos = [RetalhoLote(r.setor, r.espessura, r.descricao, r.medida, r.projeto, r.encontrado,
                            analise.indice_passagem(r))
                for r in analise.retalhos]  # mantém a ordem da análise (projeto, material)
    passagens = analise.passagens
    # mesmo texto que era colado à mão nas observações: "Retalho AMARELO: sobra do projeto X p/ Y"
    obs = "\n".join(x for x in (texto_passagens(passagens), observacoes.strip()) if x)
    return CapaLote(lote, datetime.now(), projetos, chapas, retalhos, obs, passagens)


def nome_cor_passagem(indice: int) -> str:
    return CORES_SOBRAS[indice % len(CORES_SOBRAS)][0]


def cor_passagem(indice: int) -> str:
    return CORES_SOBRAS[indice % len(CORES_SOBRAS)][1]


def texto_passagens(passagens: list[tuple[str, str]]) -> str:
    return "\n".join(f"Retalho {nome_cor_passagem(i)}: sobra do projeto {origem} p/ {destino}"
                     for i, (origem, destino) in enumerate(passagens))


# ==================================================================== layout
_COLS = 10  # grade A:J
_LARGURAS = [9, 10, 10, 11, 11, 12, 11, 11, 11, 11]  # A estreita: setor (1 a 20) e nº do projeto
# =========================================================================
#  CORES DA CAPA  -  ajuste aqui
#  Formato RRGGBB (hexadecimal, sem "#"). Depois de mudar, veja o resultado com:
#      .venv\Scripts\python -m retalhos.previa_capa
#  (gera teste_lote\PREVIA_CAPA.pdf com dados de exemplo e abre o PDF)
#
#  Dica para economizar tinta: fundos claros + texto escuro, por exemplo
COR_CABECALHO_FUNDO = "E8EEF7"
COR_CABECALHO_TITULO = "1F2937"
COR_CABECALHO_DETALHE = "4B5563"
COR_MONTAGEM_FUNDO = "DBE7FF"
COR_MONTAGEM_TEXTO = "1F2937"
# =========================================================================
# COR_CABECALHO_FUNDO = "1F2937"    # faixa do topo da frente ("LOTE 2566")
# COR_CABECALHO_TITULO = "FFFFFF"   # texto "LOTE 2566"
# COR_CABECALHO_DETALHE = "D1D5DB"  # "Emitido em ..." e o subtítulo da faixa
COR_SECAO = "2563EB"              # títulos das seções (PROJETOS, RETALHOS...) e a linha abaixo deles
COR_TABELA_CABECALHO = "DBE7FF"   # fundo do cabeçalho das tabelas e da linha TOTAL
COR_ZEBRA = "F5F7FB"              # fundo das linhas alternadas
COR_NAO_ENCONTRADO = "FDE2E2"     # retalho não encontrado no estoque
# COR_MONTAGEM_FUNDO = "2563EB"     # faixa de cada projeto no verso (controle de montagem)
# COR_MONTAGEM_TEXTO = "FFFFFF"     # texto dessa faixa
COR_TEXTO = "1F2937"              # texto normal
COR_TEXTO_SUAVE = "6B7280"        # rótulos e observações em cinza
COR_BORDA = "D0D7E2"              # linhas das tabelas
COR_PREENCHER = "9CA3AF"          # traços "____/____" para preencher à mão no verso
# cartões de resumo da frente: (fundo, cor do número) para Projetos, Chapas, Retalhos, Próx. entrega
CORES_CARTOES = [("EEF4FF", "2563EB"), ("ECFDF3", "16A34A"), ("FFF7E6", "B45309"), ("F3F4F6", "1F2937")]
# sobras que passam de um projeto do lote para o seguinte: 1ª passagem, 2ª passagem... (nome usado nas
# observações, cor de fundo da linha). Se houver mais passagens que cores, a lista recomeça.
CORES_SOBRAS = [("AMARELO", "FFF3BF"), ("VERDE", "D3F9D8"), ("AZUL", "D0EBFF"), ("LARANJA", "FFE0B2"),
                ("LILÁS", "E5DBFF"), ("ROSA", "FFDEEB"), ("CINZA", "E9ECEF")]
# =========================================================================

_FONTE = "Calibri"
_fina = Side(style="thin", color=COR_BORDA)
_BOX = Border(left=_fina, right=_fina, top=_fina, bottom=_fina)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ESQ = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

_ETAPAS = ["Início", "Corte", "Customiz.", "Coladeira", "Usinagem / Furação", "Montagem", "Painéis",
           "Embalagem", "Fim"]


@dataclass(frozen=True)
class _Densidade:
    """Alturas (pt) do layout. O compacto entra quando o conteúdo não cabe em 1 página."""

    linha: float
    cabecalho: float
    secao: float
    espaco: float
    producao: float
    fonte: float


_NORMAL = _Densidade(linha=18, cabecalho=20, secao=22, espaco=12, producao=18, fonte=10)
_COMPACTO = _Densidade(linha=14, cabecalho=16, secao=18, espaco=5, producao=16, fonte=9)

# área útil do A4 retrato com as margens usadas (pontos)
_MARGEM_LADOS, _MARGEM_TOPO = 0.4, 0.5  # polegadas
_A4_LARGURA = (8.27 - 2 * _MARGEM_LADOS) * 72
_A4_ALTURA = (11.69 - 2 * _MARGEM_TOPO) * 72
ESCALA_MINIMA_LEGIVEL = 0.70  # abaixo disso o app avisa que a impressão ficará pequena


@dataclass
class ImpressaoCapa:
    """Ajuste de página de cada face: escala estimada (1.0 = 100%) e se precisou do modo compacto."""

    escala_frente: float
    escala_verso: float
    frente_compacta: bool
    verso_compacto: bool

    @property
    def avisos(self) -> list[str]:
        return [
            f"A {face} tem muito conteúdo e será impressa reduzida a ~{escala:.0%}."
            for face, escala in (("frente", self.escala_frente), ("verso", self.escala_verso))
            if escala < ESCALA_MINIMA_LEGIVEL
        ]


def _num(valor: str) -> float:
    try:
        return float(valor)
    except ValueError:
        return float("inf")


def _col(n: int) -> str:
    return "ABCDEFGHIJ"[n - 1]


def _mm(espessura: str) -> str:
    return f"{espessura} mm" if espessura else ""


class _Folha:
    """Pequeno ajudante para escrever blocos mesclados com estilo em toda a área."""

    def __init__(self, ws: Worksheet, d: _Densidade) -> None:
        self.ws = ws
        self.d = d
        self.row = 1

    def fonte(self, size: float | None = None, color: str = COR_TEXTO, **kw: Any) -> Font:
        return Font(name=_FONTE, size=size or self.d.fonte, color=color, **kw)

    def cell(self, row: int, c1: int, c2: int, valor: Any = None, *, fonte: Font | None = None,
             fundo: str | None = None, alin: Alignment = _CENTRO, borda: Border | None = _BOX) -> None:
        ws = self.ws
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        ws.cell(row, c1, valor)
        for c in range(c1, c2 + 1):
            cel = ws.cell(row, c)
            cel.font = fonte or self.fonte()
            cel.alignment = alin
            if fundo:
                cel.fill = PatternFill("solid", fgColor=fundo)
            if borda:
                cel.border = borda

    def altura(self, row: int, pontos: float) -> None:
        self.ws.row_dimensions[row].height = pontos

    def espaco(self, pontos: float | None = None) -> None:
        self.altura(self.row, self.d.espaco if pontos is None else pontos)
        self.row += 1

    def secao(self, titulo: str, detalhe: str = "") -> None:
        r = self.row
        self.cell(r, 1, 7, titulo.upper(), fonte=self.fonte(self.d.fonte + 2, COR_SECAO, bold=True),
                  alin=Alignment(horizontal="left", vertical="bottom"), borda=None)
        self.cell(r, 8, _COLS, detalhe, fonte=self.fonte(self.d.fonte - 1, COR_TEXTO_SUAVE),
                  alin=Alignment(horizontal="right", vertical="bottom"), borda=None)
        linha = Border(bottom=Side(style="medium", color=COR_SECAO))
        for c in range(1, _COLS + 1):
            self.ws.cell(r, c).border = linha
        self.altura(r, self.d.secao)
        self.altura(r + 1, 4)
        self.row += 2

    def cabecalho(self, colunas: list[tuple[str, int, int]], altura: float | None = None) -> None:
        for titulo, c1, c2 in colunas:
            self.cell(self.row, c1, c2, titulo, fonte=self.fonte(self.d.fonte - 1, bold=True), fundo=COR_TABELA_CABECALHO)
        self.altura(self.row, altura or self.d.cabecalho)
        self.row += 1

    def linha(self, valores: list[tuple[Any, int, int]], indice: int, altura: float | None = None,
              esquerda: set[int] | None = None, fundo: str | None = None, negrito: set[int] | None = None) -> None:
        esquerda, negrito = esquerda or set(), negrito or set()
        fundo = fundo or (COR_ZEBRA if indice % 2 else None)
        for valor, c1, c2 in valores:
            self.cell(self.row, c1, c2, valor, fundo=fundo, alin=_ESQ if c1 in esquerda else _CENTRO,
                      fonte=self.fonte(bold=c1 in negrito))
        self.altura(self.row, altura or self.d.linha)
        self.row += 1

    def total(self, rotulo: str, valor: Any, ate: int) -> None:
        negrito = self.fonte(bold=True)
        self.cell(self.row, 1, ate, rotulo, fonte=negrito, fundo=COR_TABELA_CABECALHO, alin=_ESQ)
        self.cell(self.row, ate + 1, _COLS, valor, fonte=negrito, fundo=COR_TABELA_CABECALHO)
        self.altura(self.row, self.d.linha)
        self.row += 1

    def vazio(self, texto: str) -> None:
        self.cell(self.row, 1, _COLS, texto, fonte=self.fonte(self.d.fonte - 1, COR_TEXTO_SUAVE, italic=True))
        self.altura(self.row, self.d.linha)
        self.row += 1


# ------------------------------------------------------------------ medição
def _largura_pontos() -> float:
    return sum((w * 7 + 5) * 0.75 for w in _LARGURAS)  # largura de coluna do Excel -> pontos


def _altura_pontos(ws: Worksheet, ultima_linha: int) -> float:
    return sum(ws.row_dimensions[r].height or 15 for r in range(1, ultima_linha + 1))


def _escala_largura() -> float:
    return min(1.0, _A4_LARGURA / _largura_pontos())


def _escala(ws: Worksheet, ultima_linha: int) -> float:
    """Escala que o Excel usará para caber em 1 página A4 (largura e altura)."""
    return min(_escala_largura(), _A4_ALTURA / _altura_pontos(ws, ultima_linha))


# -------------------------------------------------------------------- faces
def _preparar(ws: Worksheet, capa: CapaLote, face: str) -> None:
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(_LARGURAS, start=1):
        ws.column_dimensions[_col(i)].width = w
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1  # cada face sempre em exatamente 1 página
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = _MARGEM_LADOS
    ws.page_margins.top = ws.page_margins.bottom = _MARGEM_TOPO
    ws.oddFooter.left.text = f"Lote {capa.lote}"
    ws.oddFooter.right.text = face
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8


def _frente(ws: Worksheet, capa: CapaLote, d: _Densidade) -> int:
    f = _Folha(ws, d)
    compacto = d is _COMPACTO

    # ------------------------------------------------------------ cabeçalho
    f.cell(1, 1, 6, f"LOTE {capa.lote}", fonte=f.fonte(22, COR_CABECALHO_TITULO, bold=True), fundo=COR_CABECALHO_FUNDO,
           alin=Alignment(horizontal="left", vertical="center", indent=1), borda=None)
    f.cell(1, 7, _COLS, f"Emitido em\n{capa.emitido:%d/%m/%Y %H:%M}", fonte=f.fonte(10, COR_CABECALHO_DETALHE), fundo=COR_CABECALHO_FUNDO,
           alin=Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1), borda=None)
    f.altura(1, 34 if compacto else 40)
    f.cell(2, 1, _COLS, "Capa de produção  ·  materiais, retalhos e controle do lote", fonte=f.fonte(9, COR_CABECALHO_DETALHE),
           fundo=COR_CABECALHO_FUNDO, alin=Alignment(horizontal="left", vertical="top", indent=1), borda=None)
    f.altura(2, 16)
    f.row = 3
    f.espaco()

    cartoes = [("Projetos", len(capa.projetos), 1, 2), ("Chapas inteiras", capa.total_chapas, 3, 5),
               ("Retalhos", capa.total_retalhos, 6, 8), ("Próx. entrega", capa.proxima_entrega or "—", 9, 10)]
    r = f.row
    for (titulo, valor, c1, c2), (fundo, cor) in zip(cartoes, CORES_CARTOES, strict=True):
        borda = Border(left=Side(style="thick", color=cor))
        f.cell(r, c1, c2, valor, fonte=f.fonte(13 if compacto else 16, cor, bold=True), fundo=fundo,
               alin=Alignment(horizontal="left", vertical="bottom", indent=1), borda=borda)
        f.cell(r + 1, c1, c2, titulo, fonte=f.fonte(8, COR_TEXTO_SUAVE), fundo=fundo,
               alin=Alignment(horizontal="left", vertical="top", indent=1), borda=borda)
    f.altura(r, 20 if compacto else 26)
    f.altura(r + 1, 14)
    f.row = r + 2
    f.espaco()

    # ------------------------------------------------------------- projetos
    f.secao("Projetos do lote", f"{len(capa.projetos)} projeto(s)")
    f.cabecalho([("Projeto", 1, 1), ("Pedido", 2, 2), ("Ambiente", 3, 5), ("Cliente", 6, 8), ("Entrega", 9, 10)])
    for i, p in enumerate(capa.projetos):
        f.linha([(p.projeto, 1, 1), (p.pedido, 2, 2), (p.ambiente, 3, 5), (p.cliente, 6, 8), (p.entrega, 9, 10)],
                i, esquerda={3, 6})
    if not capa.projetos:
        f.vazio("Nenhum projeto informado.")
    f.espaco()

    # ------------------------------------------------------- chapas inteiras
    # lista (material, espessura, medida, qtd): aceita qualquer espessura (6, 10, 15, 18, 25 mm...)
    f.secao("Chapas inteiras", f"{capa.total_chapas} chapa(s)")
    if capa.chapas:
        f.cabecalho([("Material", 1, 6), ("Esp.", 7, 7), ("Medida", 8, 9), ("Qtd", 10, 10)])
        ordem = sorted(capa.chapas, key=lambda c: (c.material, _num(c.espessura), c.medida))
        for i, c in enumerate(ordem):
            f.linha([(c.material, 1, 6), (_mm(c.espessura), 7, 7), (c.medida, 8, 9), (c.quantidade, 10, 10)],
                    i, esquerda={1}, negrito={10})
        f.total("TOTAL", capa.total_chapas, 9)
    else:
        f.vazio("Nenhuma chapa inteira.")
    f.espaco()

    # -------------------------------------------------------------- retalhos
    # uma linha por retalho (cada retalho é único); setor é curto (1 a 20), o espaço vai para o material
    f.secao("Retalhos", f"{capa.total_retalhos} retalho(s)")
    if capa.retalhos:
        f.cabecalho([("Setor", 1, 1), ("Esp.", 2, 2), ("Material", 3, 7), ("Medida", 8, 9), ("Projeto", 10, 10)])
        for i, r in enumerate(capa.retalhos):
            # sobra de outro projeto do lote tem a cor da passagem (vem do corte anterior, pode ainda
            # não estar no estoque); senão, vermelho quando não foi encontrado no estoque
            if r.passagem is not None:
                fundo = cor_passagem(r.passagem)
            else:
                fundo = None if r.encontrado else COR_NAO_ENCONTRADO
            f.linha([(r.setor if r.encontrado else "—", 1, 1), (_mm(r.espessura), 2, 2), (r.material, 3, 7),
                     (r.medida, 8, 9), (r.projeto, 10, 10)], i, esquerda={3}, fundo=fundo)
        # o significado de cada cor de sobra fica só nas Observações (padrão já usado na fábrica)
        legenda = f.fonte(f.d.fonte - 2, COR_TEXTO_SUAVE, italic=True)
        if any(not r.encontrado and r.passagem is None for r in capa.retalhos):
            f.cell(f.row, 1, _COLS, "Linhas em vermelho: retalho não encontrado no estoque.", fonte=legenda,
                   alin=Alignment(horizontal="left", indent=1), borda=None)
            f.altura(f.row, 13)
            f.row += 1
    else:
        f.vazio("Nenhum retalho.")
    f.espaco()

    # ----------------------------------------------------------- observações
    f.secao("Observações")
    r = f.row
    # a caixa cresce com o texto (as linhas das sobras entre projetos são preenchidas automaticamente)
    linhas_texto = len(capa.observacoes.splitlines()) if capa.observacoes else 0
    linhas_obs = max(2 if compacto else 3, linhas_texto)
    altura_obs = 14 if linhas_texto > 3 else 18
    ws.merge_cells(start_row=r, start_column=1, end_row=r + linhas_obs - 1, end_column=_COLS)
    ws.cell(r, 1, capa.observacoes or None)
    ws.cell(r, 1).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    ws.cell(r, 1).font = f.fonte()
    for rr in range(r, r + linhas_obs):
        f.altura(rr, altura_obs)
        for c in range(1, _COLS + 1):
            ws.cell(rr, c).border = _BOX
    f.row = r + linhas_obs
    f.espaco()

    # --------------------------------------------------- controle de produção
    f.secao("Controle de produção", "preencher com data / visto")
    f.cabecalho([("Projeto", 1, 1)] + [(e, i, i) for i, e in enumerate(_ETAPAS, start=2)],
                altura=24 if compacto else 26)
    for i, p in enumerate(capa.projetos or [ProjetoLote("")] * 3):
        f.linha([(p.projeto, 1, 1)] + [(None, c, c) for c in range(2, _COLS + 1)], i, altura=d.producao)
    return f.row - 1


def _verso(ws: Worksheet, capa: CapaLote, d: _Densidade) -> int:
    """Controle de montagem: um cartão por projeto, sempre numa única página (verso da folha)."""
    f = _Folha(ws, d)
    compacto = d is _COMPACTO
    f.secao(f"Controle de montagem  ·  lote {capa.lote}", f"{len(capa.projetos)} projeto(s)")
    rotulo = f.fonte(8 if compacto else 9, COR_TEXTO_SUAVE, bold=True)
    escrita = f.fonte(9, COR_PREENCHER)
    data_hora = "____ /____ /______    ____:____"
    alt = 17 if compacto else 22
    for p in capa.projetos or [ProjetoLote("")]:
        r = f.row
        partes = (f"Projeto {p.projeto}" if p.projeto else "Projeto", p.ambiente, p.cliente)
        f.cell(r, 1, _COLS, "  ·  ".join(x for x in partes if x), fonte=f.fonte(10 if compacto else 11, COR_MONTAGEM_TEXTO, bold=True),
               fundo=COR_MONTAGEM_FUNDO, alin=Alignment(horizontal="left", vertical="center", indent=1), borda=None)
        f.altura(r, 18 if compacto else 22)
        r += 1
        for esquerda, direita in (("Montador", "Resp. painéis"), ("Início", "Início"), ("Fim", "Fim")):
            f.cell(r, 1, 1, esquerda, fonte=rotulo, alin=_ESQ)
            f.cell(r, 2, 5, None if esquerda == "Montador" else data_hora, alin=_ESQ, fonte=escrita)
            f.cell(r, 6, 6, direita, fonte=rotulo, alin=_ESQ)
            f.cell(r, 7, _COLS, None if direita.startswith("Resp") else data_hora, alin=_ESQ, fonte=escrita)
            f.altura(r, alt)
            r += 1
        f.cell(r, 1, 5, "Pendências montagem", fonte=rotulo, alin=_ESQ, fundo=COR_ZEBRA)
        f.cell(r, 6, _COLS, "Pendências painéis", fonte=rotulo, alin=_ESQ, fundo=COR_ZEBRA)
        f.altura(r, 14 if compacto else 18)
        r += 1
        for _ in range(1 if compacto else 2):
            f.cell(r, 1, 5, None)
            f.cell(r, 6, _COLS, None)
            f.altura(r, alt + 3)
            r += 1
        f.row = r
        f.espaco(6 if compacto else 12)
    return f.row - 1


def _montar_face(wb: Workbook, titulo: str, capa: CapaLote,
                 construir: Callable[[Worksheet, CapaLote, _Densidade], int]) -> tuple[Worksheet, float, bool]:
    """Monta a face no modo normal; se não couber em 1 página sem reduzir, refaz no compacto."""
    for d in (_NORMAL, _COMPACTO):
        ws = wb.create_sheet(titulo)
        _preparar(ws, capa, titulo.lower())
        ultima = construir(ws, capa, d)
        ws.print_area = f"A1:{_col(_COLS)}{ultima}"
        escala = _escala(ws, ultima)
        if d is _COMPACTO or escala >= _escala_largura() - 1e-6:
            return ws, escala, d is _COMPACTO
        wb.remove(ws)
    raise AssertionError("inalcançável")


def montar_capa(capa: CapaLote) -> tuple[Workbook, ImpressaoCapa]:
    """Pasta com as abas ``Frente`` e ``Verso``, cada uma ajustada para 1 página A4."""
    wb = Workbook()
    wb.remove(wb.active)
    frente, escala_f, compacta_f = _montar_face(wb, "Frente", capa, _frente)
    _, escala_v, compacto_v = _montar_face(wb, "Verso", capa, _verso)
    wb.active = wb.index(frente)
    return wb, ImpressaoCapa(escala_f, escala_v, compacta_f, compacto_v)


def avaliar_impressao(capa: CapaLote) -> ImpressaoCapa:
    """Como a capa ficará na impressão (sem gravar arquivo)."""
    return montar_capa(capa)[1]


def gerar_capa(capa: CapaLote, destino: str | Path) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb, _ = montar_capa(capa)
    wb.save(destino)
    return destino


# ====================================================================== PDF
_PS_EXPORTAR = r"""
$ErrorActionPreference = 'Stop'
$x = New-Object -ComObject Excel.Application
$x.DisplayAlerts = $false
try {
    $wb = $x.Workbooks.Open($env:CAPA_XLSX, 0, $true)
    $wb.ExportAsFixedFormat(0, $env:CAPA_PDF)
    $wb.Close($false)
} finally { $x.Quit() }
"""


def exportar_pdf(xlsx: str | Path, pdf: str | Path | None = None) -> Path:
    """PDF para impressão frente e verso (página 1 = frente, página 2 = verso), gerado pelo Excel instalado.

    Lança ``RuntimeError`` se o Excel não estiver disponível ou a exportação falhar.
    """
    xlsx = Path(xlsx).resolve()
    pdf = Path(pdf).resolve() if pdf else xlsx.with_suffix(".pdf")
    env = dict(os.environ, CAPA_XLSX=str(xlsx), CAPA_PDF=str(pdf))  # caminhos via ambiente: sem problema de aspas
    try:
        proc = subprocess.run(  # noqa: S603
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", _PS_EXPORTAR],
            env=env, capture_output=True, text=True, timeout=180,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise RuntimeError(f"Não foi possível gerar o PDF: {exc}") from exc
    if proc.returncode != 0 or not pdf.is_file():
        detalhe = (proc.stderr or proc.stdout or "").strip()[:500]
        raise RuntimeError(f"Não foi possível gerar o PDF pelo Excel.\n{detalhe}")
    return pdf


__all__ = [
    "CapaLote", "ChapaLote", "ImpressaoCapa", "ProjetoLote", "RetalhoLote", "avaliar_impressao",
    "capa_de_analise", "exportar_pdf", "gerar_capa", "ler_conf_folha_lote", "montar_capa",
]

